# -*- coding: utf-8 -*-
"""
CHUYỂN ĐỔI FILE (giữ nguyên bố cục, KHÔNG sửa file gốc)
  - PDF  -> Excel : bảng thành ô Excel (viền, gộp ô, màu nền), chữ đặt đúng toạ độ,
                    chữ ký / con dấu / logo thành ảnh trong suốt đặt đúng chỗ.
  - PDF  -> Word  : dùng core (_docx_build) – nền cố định + chữ đúng toạ độ.
  - Word / Excel -> PDF : dùng Microsoft Office (hoặc WPS, LibreOffice) có trên máy,
                    chuyển từ BẢN SAO TẠM nên file gốc không bị đụng tới.
"""
import os
import io
import sys
import shutil
import bisect
import tempfile
import subprocess

import core
from core import fitz

WORD_EXT = (".doc", ".docx", ".docm", ".rtf", ".odt", ".dot", ".dotx")
EXCEL_EXT = (".xls", ".xlsx", ".xlsm", ".xlsb", ".ods", ".csv")


# =============================================================================
#  TIỆN ÍCH
# =============================================================================
def unique_path(path):
    """Không bao giờ ghi đè: thêm (1), (2)... nếu file đã có."""
    if not os.path.exists(path):
        return path
    base, ext = os.path.splitext(path)
    i = 1
    while os.path.exists("%s (%d)%s" % (base, i, ext)):
        i += 1
    return "%s (%d)%s" % (base, i, ext)


def _snap(values, tol):
    vals = sorted(values)
    out = []
    for v in vals:
        if not out or v - out[-1] > tol:
            out.append(v)
    return out


def _nearest(arr, v):
    i = bisect.bisect_left(arr, v)
    if i <= 0:
        return 0
    if i >= len(arr):
        return len(arr) - 1
    return i if arr[i] - v < v - arr[i - 1] else i - 1


def _hex(c):
    return "%06X" % (c & 0xFFFFFF)


# Excel: độ rộng cột tính theo "ký tự" của font chuẩn Calibri 11 (7 px / ký tự, 96 dpi)
def _pt_to_colwidth(pt):
    px = pt * 96.0 / 72.0
    if px < 12:
        return round(px / 12.0, 4)
    return round((px - 5) / 7.0, 4)


def _colwidth_to_pt(w):
    px = w * 12 if w < 1 else int(w * 7 + 5)
    return px * 72.0 / 96.0


# =============================================================================
#  PDF -> EXCEL
# =============================================================================
def _page_lines(page):
    """Các dòng chữ NGANG (theo hướng hiển thị), toạ độ hiển thị."""
    rot = page.rotation_matrix
    rot0 = fitz.Matrix(rot.a, rot.b, rot.c, rot.d, 0, 0)
    dd = page.get_text("rawdict", flags=fitz.TEXT_PRESERVE_WHITESPACE | fitz.TEXT_PRESERVE_LIGATURES)
    out = []
    for b in dd["blocks"]:
        if b.get("type", 0) != 0:
            continue
        for ln in b["lines"]:
            ink = fitz.Rect()
            for s in ln["spans"]:
                chars = s.pop("chars", [])
                s["text"] = "".join(c["c"] for c in chars)
                for c in chars:
                    if core.norm_text(c["c"]).strip():
                        ink |= fitz.Rect(c["bbox"])
            spans = [s for s in ln["spans"] if core.norm_text(s["text"]).strip()]
            if not spans:
                continue
            if all(s.get("alpha", 255) == 0 for s in spans):
                continue
            dv = fitz.Point(ln.get("dir", (1, 0))) * rot0
            if abs(dv.x - 1) > 1e-3:
                continue
            # khung chữ THẬT (bỏ dấu cách thừa cuối dòng), cao theo dòng
            lb = fitz.Rect(ln["bbox"])
            if not ink.is_empty:
                lb = fitz.Rect(ink.x0, lb.y0, ink.x1, lb.y1)
            bbox = lb * rot
            out.append({"bbox": bbox, "spans": spans, "raw": ln})
    return out


def _line_text(ln):
    return core.norm_text("".join(s["text"] for s in ln["raw"]["spans"])).strip()


def _font_of(spans):
    from openpyxl.styles import Font
    s = max(spans, key=lambda s: len(s["text"].strip()))
    fname = s.get("font", "")
    fl = s.get("flags", 0)
    return Font(name=core.clean_font_name(fname), size=round(s["size"] * 2) / 2,
                bold=bool(fl & 16) or "bold" in fname.lower(),
                italic=bool(fl & 2) or "italic" in fname.lower(),
                color=_hex(s.get("color", 0)))


def _graphics_layer(doc, pno, lines, tables, dpi):
    """Ảnh trong suốt chứa mọi thứ KHÔNG phải chữ ngang và KHÔNG phải khung bảng:
    chữ ký, con dấu, logo, đường kẻ trang trí... Trả về list (rect_pt, png_bytes)."""
    import numpy as np
    tmp = fitz.open()
    tmp.insert_pdf(doc, from_page=pno, to_page=pno)
    tp = tmp[0]
    derot = tp.derotation_matrix
    for ln in lines:
        for s in ln["raw"]["spans"]:
            if s["text"].strip():
                tp.add_redact_annot(core.PdfDocument._shrunk(s["bbox"], tp, 0.15), fill=False)
    core.apply_text_only_redactions(tp)
    # bỏ khung bảng + nền ô (đã chuyển thành viền / màu ô Excel)
    covered = getattr(fitz, "PDF_REDACT_LINE_ART_REMOVE_IF_COVERED", 1)
    for t in tables:
        r = fitz.Rect(t.bbox)
        r = fitz.Rect(r.x0 - 1.5, r.y0 - 1.5, r.x1 + 1.5, r.y1 + 1.5) * derot
        tp.add_redact_annot(r, fill=False)
    if tables:
        try:
            tp.apply_redactions(images=core.REDACT_KW["images"], graphics=covered,
                                text=core.REDACT_KW["text"])
        except TypeError:
            pass
    pm = tp.get_pixmap(dpi=dpi, alpha=True)
    if pm.n != 4:
        return []
    arr = np.frombuffer(pm.samples, dtype=np.uint8).reshape(pm.height, pm.width, 4)
    mask = (arr[:, :, 3] > 8).astype(np.uint8)
    if not mask.any():
        return []
    scale = 72.0 / dpi
    pieces = []
    try:
        import cv2
        k = max(3, int(dpi / 18))
        dil = cv2.dilate(mask, np.ones((k, k), np.uint8))
        n, _, stats, _ = cv2.connectedComponentsWithStats(dil, connectivity=8)
        boxes = [tuple(stats[i][:4]) for i in range(1, n)]
    except Exception:
        ys, xs = np.nonzero(mask)
        boxes = [(xs.min(), ys.min(), xs.max() - xs.min() + 1, ys.max() - ys.min() + 1)]
    for (x, y, w, h) in boxes:
        sub = mask[y:y + h, x:x + w]
        if sub.sum() < 4:
            continue
        ys, xs = np.nonzero(sub)
        x0, y0 = x + xs.min(), y + ys.min()
        x1, y1 = x + xs.max() + 1, y + ys.max() + 1
        rgba = arr[y0:y1, x0:x1].copy()
        from PIL import Image
        im = Image.fromarray(rgba, "RGBA")
        bio = io.BytesIO()
        im.save(bio, "PNG")
        pieces.append((fitz.Rect(x0 * scale, y0 * scale, x1 * scale, y1 * scale), bio.getvalue()))
    return pieces


def pdf_to_excel(doc, out, dpi=200, progress=None):
    """doc: fitz.Document. Mỗi trang PDF -> 1 trang tính, bố cục như bản gốc."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Side, PatternFill
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties

    EMU = 12700
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="000000")
    n = doc.page_count
    for pno in range(n):
        page = doc[pno]
        W, H = page.rect.width, page.rect.height
        ws = wb.create_sheet("Trang %d" % (pno + 1))
        ws.sheet_view.showGridLines = False

        lines = _page_lines(page)
        try:
            tables = [t for t in page.find_tables().tables if t.row_count and t.col_count]
        except Exception:
            tables = []
        # các ô bảng (toạ độ hiển thị)
        rot = page.rotation_matrix
        cells = []  # (rect, table_index)
        for ti, t in enumerate(tables):
            for c in t.cells:
                if c:
                    cells.append((fitz.Rect(c) * rot if page.rotation else fitz.Rect(c), ti))
        tab_rects = [fitz.Rect(t.bbox) * rot if page.rotation else fitz.Rect(t.bbox) for t in tables]

        # phân dòng chữ: trong ô bảng / tự do
        cell_lines = {i: [] for i in range(len(cells))}
        free = []
        for ln in lines:
            c = (ln["bbox"].tl + ln["bbox"].br) / 2
            hit = None
            for i, (r, _) in enumerate(cells):
                if r.contains(c):
                    if hit is None or r.get_area() < cells[hit][0].get_area():
                        hit = i
            if hit is None:
                free.append(ln)
            else:
                cell_lines[hit].append(ln)

        # ---- lưới cột / hàng theo toạ độ PDF
        xs = [0.0, W]
        ys = [0.0, H]
        for r, _ in cells:
            xs += [r.x0, r.x1]
            ys += [r.y0, r.y1]
        # chữ tự do: gom theo hàng (cùng đường chân chữ) -> mỗi hàng chữ = 1 hàng Excel
        groups = []
        for ln in sorted(free, key=lambda l: l["raw"]["spans"][0]["origin"][1]):
            by = (fitz.Point(ln["raw"]["spans"][0]["origin"]) * rot).y
            if groups and abs(groups[-1]["by"] - by) < 2.5:
                g = groups[-1]
                g["top"] = min(g["top"], ln["bbox"].y0)
                g["bot"] = max(g["bot"], ln["bbox"].y1)
                g["lines"].append(ln)
            else:
                groups.append({"by": by, "top": ln["bbox"].y0, "bot": ln["bbox"].y1, "lines": [ln]})
        for g in groups:
            ys += [g["top"], g["bot"]]
            for ln in g["lines"]:
                ln["grow"] = g
                xs.append(ln["bbox"].x0)
        xs = _snap([min(max(v, 0), W) for v in xs], 1.2)
        ys = _snap([min(max(v, 0), H) for v in ys], 1.0)
        if xs[-1] < W - 1.2:
            xs.append(W)
        if ys[-1] < H - 1.0:
            ys.append(H)
        ncol, nrow = len(xs) - 1, len(ys) - 1
        for ci in range(ncol):
            ws.column_dimensions[get_column_letter(ci + 1)].width = _pt_to_colwidth(xs[ci + 1] - xs[ci])
        for ri in range(nrow):
            ws.row_dimensions[ri + 1].height = round(ys[ri + 1] - ys[ri], 2)
        # toạ độ thực của cột sau khi Excel làm tròn độ rộng (để đặt ảnh chính xác)
        xs_real = [0.0]
        for ci in range(ncol):
            xs_real.append(xs_real[-1] + _colwidth_to_pt(ws.column_dimensions[get_column_letter(ci + 1)].width))

        used = [[False] * ncol for _ in range(nrow)]

        def span(r):
            r0, r1 = _nearest(ys, r.y0), _nearest(ys, r.y1)
            c0, c1 = _nearest(xs, r.x0), _nearest(xs, r.x1)
            return r0, max(r1, r0 + 1), c0, max(c1, c0 + 1)

        def free_area(r0, r1, c0, c1):
            return all(not used[i][j] for i in range(r0, r1) for j in range(c0, c1))

        def occupy(r0, r1, c0, c1):
            for i in range(r0, r1):
                for j in range(c0, c1):
                    used[i][j] = True

        # màu nền ô bảng
        fills = []
        try:
            for dr in page.get_drawings():
                if dr.get("fill") and dr.get("rect") is not None:
                    fills.append((fitz.Rect(dr["rect"]) * rot if page.rotation else fitz.Rect(dr["rect"]), dr["fill"]))
        except Exception:
            pass

        # ---- ô bảng
        for i, (r, ti) in enumerate(cells):
            r0, r1, c0, c1 = span(r)
            occupy(r0, r1, c0, c1)
            if r1 - r0 > 1 or c1 - c0 > 1:
                ws.merge_cells(start_row=r0 + 1, start_column=c0 + 1, end_row=r1, end_column=c1)
            for rr in range(r0, r1):
                for cc in range(c0, c1):
                    ws.cell(rr + 1, cc + 1).border = Border(
                        left=thin if cc == c0 else None, right=thin if cc == c1 - 1 else None,
                        top=thin if rr == r0 else None, bottom=thin if rr == r1 - 1 else None)
            cell = ws.cell(r0 + 1, c0 + 1)
            center = (r.tl + r.br) / 2
            for fr, col in fills:
                if fr.contains(center) and fr.get_area() <= tab_rects[ti].get_area() * 1.01 and col != (1, 1, 1) and len(col) == 3:
                    hx = "%02X%02X%02X" % tuple(int(round(v * 255)) for v in col)
                    for rr in range(r0, r1):
                        for cc in range(c0, c1):
                            ws.cell(rr + 1, cc + 1).fill = PatternFill("solid", fgColor=hx)
                    break
            lns = sorted(cell_lines[i], key=lambda l: (round(l["bbox"].y0), l["bbox"].x0))
            if not lns:
                continue
            # ghép các dòng cùng hàng
            rows_txt, last_y = [], None
            for l in lns:
                t = _line_text(l)
                if last_y is not None and abs(l["bbox"].y0 - last_y) < 2:
                    rows_txt[-1] += " " + t
                else:
                    rows_txt.append(t)
                last_y = l["bbox"].y0
            text = "\n".join(rows_txt)
            cell.value = _num(text)
            cell.font = _font_of([s for l in lns for s in l["spans"]])
            tb = fitz.Rect(lns[0]["bbox"])
            for l in lns[1:]:
                tb |= l["bbox"]
            gl, gr = tb.x0 - r.x0, r.x1 - tb.x1
            if abs(gl - gr) < max(3, r.width * 0.08):
                ha = "center"
            elif gr < gl:
                ha = "right"
            else:
                ha = "left"
            gt, gb = tb.y0 - r.y0, r.y1 - tb.y1
            va = "center" if abs(gt - gb) < max(3, r.height * 0.15) else ("top" if gt < gb else "bottom")
            cell.alignment = Alignment(horizontal=ha, vertical=va, wrap_text=len(rows_txt) > 1,
                                       indent=1 if (ha == "left" and gl > 3) or (ha == "right" and gr > 3) else 0)

        # ---- chữ tự do: đặt đúng toạ độ, không gộp ô -> chữ tràn sang ô trống như Excel thường
        for ln in sorted(free, key=lambda l: (l["bbox"].y0, l["bbox"].x0)):
            g = ln["grow"]
            r0 = _nearest(ys, g["top"])
            r1 = max(_nearest(ys, g["bot"]), r0 + 1)
            c0 = _nearest(xs, ln["bbox"].x0)
            if used[r0][c0]:
                cell = ws.cell(r0 + 1, c0 + 1)
                cell.value = ("%s %s" % (cell.value or "", _line_text(ln))).strip()
                continue
            if r1 - r0 > 1 and free_area(r0, r1, c0, c0 + 1):
                ws.merge_cells(start_row=r0 + 1, start_column=c0 + 1, end_row=r1, end_column=c0 + 1)
                occupy(r0, r1, c0, c0 + 1)
            else:
                occupy(r0, r0 + 1, c0, c0 + 1)
            cell = ws.cell(r0 + 1, c0 + 1)
            cell.value = _line_text(ln)
            cell.data_type = "s"
            cell.font = _font_of(ln["spans"])
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)

        # ---- chữ ký, con dấu, logo... thành ảnh đặt đúng chỗ
        max_x = max([r.x1 for r, _ in cells] + [l["bbox"].x1 for l in lines] + [1.0])
        max_y = max([r.y1 for r, _ in cells] + [l["bbox"].y1 for l in lines] + [1.0])
        for rect, png in _graphics_layer(doc, pno, lines, tables, dpi):
            max_x, max_y = max(max_x, rect.x1), max(max_y, rect.y1)
            img = XLImage(io.BytesIO(png))
            ci = max(0, bisect.bisect_right(xs, rect.x0) - 1)
            ri = max(0, bisect.bisect_right(ys, rect.y0) - 1)
            ci = min(ci, ncol - 1)
            ri = min(ri, nrow - 1)
            # lệch trong ô: theo toạ độ thực của cột (sau khi Excel làm tròn)
            col_off = (rect.x0 - xs[ci]) * (xs_real[ci + 1] - xs_real[ci]) / max(xs[ci + 1] - xs[ci], 0.01)
            row_off = rect.y0 - ys[ri]
            sx = (xs_real[-1] / W) if W else 1
            img.anchor = OneCellAnchor(
                _from=AnchorMarker(col=ci, colOff=int(col_off * EMU), row=ri, rowOff=int(row_off * EMU)),
                ext=XDRPositiveSize2D(int(rect.width * sx * EMU), int(rect.height * EMU)))
            img.width, img.height = rect.width * sx * 96 / 72, rect.height * 96 / 72
            ws.add_image(img)

        # ---- in ấn: đúng khổ giấy, lề 0, vừa 1 trang
        ws.page_setup.orientation = "landscape" if W > H else "portrait"
        a4 = (abs(min(W, H) - 595) < 8 and abs(max(W, H) - 842) < 8)
        ws.page_setup.paperSize = ws.PAPERSIZE_A4 if a4 else ws.PAPERSIZE_LETTER
        ws.page_margins.left = ws.page_margins.right = 0
        ws.page_margins.top = ws.page_margins.bottom = 0
        ws.page_margins.header = ws.page_margins.footer = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        # vùng in: tới hết nội dung (bỏ khoảng trắng cuối trang) -> in đúng tỉ lệ 100%,
        # không bị thu nhỏ do Excel làm tròn chiều cao hàng / độ rộng cột
        last_c = min(ncol, bisect.bisect_left(xs, max_x - 0.5) + 1)
        last_r = min(nrow, bisect.bisect_left(ys, max_y - 0.5) + 1)
        last_c = min(ncol, last_c + 1)
        last_r = min(nrow, last_r + 1)
        ws.print_area = "A1:%s%d" % (get_column_letter(last_c), last_r)
        ws.print_options.horizontalCentered = False
        if progress:
            progress(pno + 1, n)
    wb.save(out)
    return out


def _num(text):
    """Giữ nguyên dạng chữ (vd '02', '01/01/1990' không bị Excel đổi)."""
    return text


# =============================================================================
#  WORD / EXCEL -> PDF
# =============================================================================
def find_soffice():
    c = shutil.which("soffice") or shutil.which("libreoffice")
    if c:
        return c
    for base in (os.environ.get("ProgramFiles", r"C:\Program Files"),
                 os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")):
        for name in ("LibreOffice", "LibreOffice 7", "LibreOffice 24", "OpenOffice 4"):
            p = os.path.join(base, name, "program", "soffice.exe")
            if os.path.isfile(p):
                return p
    return None


def _com_word(src, out, progid):
    import win32com.client
    app = win32com.client.DispatchEx(progid)
    try:
        app.Visible = False
    except Exception:
        pass
    try:
        app.DisplayAlerts = 0
    except Exception:
        pass
    doc = None
    try:
        doc = app.Documents.Open(src, False, True, False)  # ConfirmConversions, ReadOnly, AddToRecent
        try:
            doc.ExportAsFixedFormat(out, 17)  # wdExportFormatPDF
        except Exception:
            doc.SaveAs2(out, 17) if hasattr(doc, "SaveAs2") else doc.SaveAs(out, 17)
    finally:
        try:
            if doc is not None:
                doc.Close(False)
        finally:
            try:
                app.Quit()
            except Exception:
                pass


def _com_excel(src, out, progid):
    import win32com.client
    app = win32com.client.DispatchEx(progid)
    try:
        app.Visible = False
        app.DisplayAlerts = False
    except Exception:
        pass
    wb = None
    try:
        wb = app.Workbooks.Open(src, 0, True)  # UpdateLinks=0, ReadOnly=True
        wb.ExportAsFixedFormat(0, out)         # xlTypePDF
    finally:
        try:
            if wb is not None:
                wb.Close(False)
        finally:
            try:
                app.Quit()
            except Exception:
                pass


def office_to_pdf(src, out, log=None):
    """Chuyển Word/Excel -> PDF. Làm trên BẢN SAO TẠM -> file gốc không đổi.
    Trả về tên chương trình đã dùng."""
    log = log or (lambda m: None)
    ext = os.path.splitext(src)[1].lower()
    kind = "word" if ext in WORD_EXT else "excel" if ext in EXCEL_EXT else None
    if kind is None:
        raise ValueError("Không hỗ trợ định dạng %s" % ext)
    work = tempfile.mkdtemp(prefix="suapdf_")
    try:
        tmp_src = os.path.join(work, "nguon" + ext)
        shutil.copy2(src, tmp_src)
        tmp_out = os.path.join(work, "ketqua.pdf")
        errors = []
        if sys.platform.startswith("win"):
            try:
                import pythoncom
                pythoncom.CoInitialize()
            except Exception:
                pass
            engines = ([("Microsoft Word", "Word.Application"), ("WPS Writer", "KWps.Application"),
                        ("WPS Writer", "wps.Application")] if kind == "word" else
                       [("Microsoft Excel", "Excel.Application"), ("WPS Spreadsheets", "KET.Application"),
                        ("WPS Spreadsheets", "et.Application")])
            for name, progid in engines:
                try:
                    (_com_word if kind == "word" else _com_excel)(tmp_src, tmp_out, progid)
                    if os.path.isfile(tmp_out) and os.path.getsize(tmp_out) > 0:
                        shutil.move(tmp_out, out)
                        return name
                except Exception as ex:
                    errors.append("%s: %s" % (name, str(ex).splitlines()[0] if str(ex) else type(ex).__name__))
        so = find_soffice()
        if so:
            prof = "file:///" + os.path.join(work, "profile").replace("\\", "/").lstrip("/")
            cmd = [so, "-env:UserInstallation=" + prof, "--headless", "--norestore",
                   "--convert-to", "pdf", "--outdir", work, tmp_src]
            flags = 0x08000000 if sys.platform.startswith("win") else 0  # CREATE_NO_WINDOW
            subprocess.run(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                           timeout=300, creationflags=flags)
            p = os.path.join(work, "nguon.pdf")
            if os.path.isfile(p) and os.path.getsize(p) > 0:
                shutil.move(p, out)
                return "LibreOffice"
            errors.append("LibreOffice: không tạo được PDF")
        raise RuntimeError(
            "Không chuyển được sang PDF.\n"
            "Máy cần cài Microsoft Office (Word/Excel), WPS Office hoặc LibreOffice "
            "để giữ đúng 100% bố cục khi chuyển Word/Excel sang PDF.\n\n" + "\n".join(errors))
    finally:
        shutil.rmtree(work, ignore_errors=True)


# =============================================================================
#  HÀM CHUNG CHO GIAO DIỆN
# =============================================================================
MODES = [
    ("pdf2word", "PDF → Word (.docx)", (".pdf",), ".docx"),
    ("pdf2excel", "PDF → Excel (.xlsx)", (".pdf",), ".xlsx"),
    ("word2pdf", "Word → PDF", WORD_EXT, ".pdf"),
    ("excel2pdf", "Excel → PDF", EXCEL_EXT, ".pdf"),
]


def detect_engines(kind):
    """Chương trình có trên máy để chuyển Word/Excel -> PDF (kind: 'word' | 'excel')."""
    found = []
    if sys.platform.startswith("win"):
        try:
            import winreg
            progs = ([("Microsoft Word", "Word.Application"), ("WPS Writer", "KWps.Application")]
                     if kind == "word" else
                     [("Microsoft Excel", "Excel.Application"), ("WPS Spreadsheets", "KET.Application")])
            for name, pid in progs:
                try:
                    winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, pid + "\\CLSID"))
                    found.append(name)
                except OSError:
                    pass
        except Exception:
            pass
    if find_soffice():
        found.append("LibreOffice")
    return found


# =============================================================================
#  TRANG TRỐNG
# =============================================================================
def is_blank_page(page):
    """Trang trống: không có chữ / hình / đường kẻ nào ở phần thân trang
    (bỏ qua vùng đầu trang & chân trang – vd trang trống chỉ còn số trang)."""
    # còn BẤT KỲ chữ nào ở phần thân trang (kể cả "./.") -> không phải trang trống
    try:
        H = page.rect.height
        for w in page.get_text("words"):
            r = fitz.Rect(w[:4]) * page.rotation_matrix
            cy = (r.y0 + r.y1) / 2
            if w[4].strip() and H * 0.09 <= cy <= H * 0.91:
                return False
    except Exception:
        return False
    try:
        import numpy as np
        pm = page.get_pixmap(dpi=40, colorspace=fitz.csGRAY, alpha=False, annots=True)
        arr = np.frombuffer(pm.samples, dtype=np.uint8).reshape(pm.height, pm.stride)[:, :pm.width]
        band = int(pm.height * 0.09)
        body = arr[band:pm.height - band]
        ink = int((body < 200).sum())
        return ink <= max(2, int(body.size * 0.0002))
    except Exception:
        txt = page.get_text("text").strip()
        return not txt and not page.get_images() and not page.get_drawings()


def drop_blank_pages(doc):
    """Xoá trang trống trong doc (luôn giữ lại ít nhất 1 trang). Trả về số trang đã xoá."""
    blank = [i for i in range(doc.page_count) if is_blank_page(doc[i])]
    if len(blank) >= doc.page_count:
        blank = blank[1:]
    for i in reversed(blank):
        doc.delete_page(i)
    return len(blank)


def convert_file(mode, src, out_dir=None, word_style="layout", dpi=200, log=None, doc=None,
                 remove_blank=True, stats=None):
    """Chuyển 1 file. Trả về (đường dẫn file kết quả, chương trình đã dùng). File gốc chỉ được ĐỌC.
    doc: (tuỳ chọn) tài liệu PDF đang mở trong phần mềm – gồm cả phần đã sửa chưa lưu.
    remove_blank: tự xoá trang trống (PDF nguồn khi chuyển sang Word/Excel; PDF kết quả khi Word/Excel -> PDF).
    stats: dict (tuỳ chọn) nhận thêm thông tin, vd stats['blank'] = số trang trống đã xoá."""
    stats = stats if stats is not None else {}
    stats["blank"] = 0
    ext = dict((m[0], m[3]) for m in MODES)[mode]
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    base = os.path.splitext(os.path.basename(src))[0]
    out = unique_path(os.path.join(out_dir or os.path.dirname(src), base + ext))
    if os.path.abspath(out) == os.path.abspath(src):
        out = unique_path(os.path.join(out_dir or os.path.dirname(src), base + "_chuyen" + ext))
    if mode in ("pdf2word", "pdf2excel"):
        if doc is None:
            with open(src, "rb") as f:
                data = f.read()
            doc = fitz.open(stream=data, filetype="pdf")
        if doc.needs_pass:
            raise RuntimeError("File PDF có mật khẩu – hãy mở bằng phần mềm rồi Lưu dạng Word/Excel.")
        if remove_blank and doc.page_count > 1:
            # làm trên BẢN SAO trong bộ nhớ – không đụng file gốc / tài liệu đang mở
            doc = fitz.open(stream=doc.tobytes(), filetype="pdf")
            stats["blank"] = drop_blank_pages(doc)
        if mode == "pdf2excel":
            pdf_to_excel(doc, out, dpi)
        elif word_style == "flow":
            pd = core.PdfDocument.__new__(core.PdfDocument)
            pd.doc = doc
            core.PdfDocument.save_docx_flow(pd, out)
        else:
            import wordnative
            wordnative.pdf_to_docx(doc, out, dpi)
        return out, "PyMuPDF"
    engine = office_to_pdf(src, out, log)
    if remove_blank:
        try:
            d = fitz.open(out)
            n = drop_blank_pages(d) if d.page_count > 1 else 0
            if n:
                tmp = out + ".tmp"
                d.save(tmp, garbage=3, deflate=True)
                d.close()
                os.replace(tmp, out)
            else:
                d.close()
            stats["blank"] = n
        except Exception:
            pass
    return out, engine
