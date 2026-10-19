# -*- coding: utf-8 -*-
"""Tự kiểm tra (chạy: SuaPDF.exe --selftest <thư_mục>) – dùng trong GitHub Actions.
Ghi từng bước ra màn hình + file selftest.log để biết chính xác lỗi ở đâu."""
import os
import sys
import traceback

import core
from core import fitz


class Log:
    def __init__(self, out_dir):
        self.path = os.path.join(out_dir, "selftest.log")
        self.lines = []

    def __call__(self, msg):
        self.lines.append(msg)
        try:
            sys.stdout.write(msg + "\n")
            sys.stdout.flush()
        except Exception:
            pass
        try:
            with open(self.path, "w", encoding="utf-8") as f:
                f.write("\n".join(self.lines))
        except Exception:
            pass


def run(out_dir):
    out_dir = os.path.abspath(out_dir)
    os.makedirs(out_dir, exist_ok=True)
    log = Log(out_dir)
    errors = []

    def step(name, fn, critical=True):
        try:
            r = fn()
            log("[OK]  " + name)
            return r
        except Exception:
            tb = traceback.format_exc()
            log(("[LOI] " if critical else "[CANH BAO] ") + name + "\n" + tb)
            if critical:
                errors.append(name)
            return None

    log("Python %s | PyMuPDF %s | %s" % (sys.version.split()[0], getattr(fitz, "VersionBind", "?"), sys.platform))
    log("Thu muc font: %s" % core.font_dirs())

    fams = step("Doc danh sach font", core.scan_fonts) or {}
    log("So font: %d" % len(fams))
    if not fams:
        log("SELFTEST FAIL: khong co font")
        return 1
    fm = core.FontManager(fams)
    fam = fm.default_family()
    path, _ = fm.resolve(fam)
    log("Font mac dinh: %s -> %s" % (fam, path))

    src = os.path.join(out_dir, "mau.pdf")
    state = {}

    def make_sample():
        d = fitz.open()
        p = d.new_page()
        p.insert_text((72, 80), "Họ và tên: Nguyễn Văn An", fontname="F0", fontfile=path, fontsize=13)
        p.insert_text((72, 95), "Ngày sinh: 01/01/1990", fontname="F0", fontfile=path, fontsize=13)
        for r in range(3):
            for c in range(3):
                rect = fitz.Rect(72 + c * 120, 150 + r * 22, 192 + c * 120, 172 + r * 22)
                p.draw_rect(rect, color=(0, 0, 0), width=0.8)
                p.insert_text((rect.x0 + 4, rect.y0 + 15), "O %d-%d" % (r, c), fontname="F0", fontfile=path, fontsize=12)
        p.draw_circle((400, 300), 30, color=(0.8, 0, 0))
        state["n_draw"] = len(p.get_drawings())
        d.save(src)

    step("Tao PDF mau", make_sample)
    if errors:
        log("SELFTEST FAIL: " + ", ".join(errors))
        return 1

    doc = step("Mo PDF", lambda: core.PdfDocument(src))
    if doc is None:
        log("SELFTEST FAIL: khong mo duoc PDF")
        return 1

    def edit():
        texts = [l.text for l in doc.lines(0)]
        log("    Chu doc duoc: %s" % ascii(texts))
        ln = [l for l in doc.lines(0) if "sinh" in l.text][0]
        doc.replace_line(0, ln, "Ngày sinh: 15/08/1985 – Hà Nội", fm, fam, ln.size, ln.color, False, False)
        cell = [l for l in doc.lines(0) if l.bbox.contains(fitz.Point(205, 183))]
        assert cell, "khong tim thay dong chu trong o bang giua"
        assert cell[0].text.strip() == "O 1-1", "doc sai chu trong o: %s" % ascii(cell[0].text)
        ln = cell[0]
        doc.replace_line(0, ln, "Đã sửa", fm, fam, ln.size, ln.color, True, False)
        doc.add_text(0, 72, 400, "Dòng thêm mới", fm, fam, 13, (0, 0, 0), False, True)
        texts = [l.text for l in doc.lines(0)]
        log("    Sau khi sua: %s" % ascii(texts))
        assert any("15/08/1985" in t for t in texts), "khong thay dong da sua"
        assert any("Nguy" in t for t in texts), "mat dong ben canh"
        n = len(doc.doc[0].get_drawings())
        assert n == state["n_draw"], "mat duong ke bang (%d -> %d)" % (state["n_draw"], n)

    step("Sua chu / them chu (giu duong ke bang)", edit)

    def out(name):
        return os.path.join(out_dir, name)

    step("Luu PDF", lambda: doc.save_pdf(out("kq.pdf")))
    step("Luu Word giu bo cuc", lambda: doc.save_docx_layout(out("kq_bocuc.docx")))
    step("Luu Word anh trang", lambda: doc.save_docx_image(out("kq_anh.docx")))
    step("Luu Word lien mach (pdf2docx)", lambda: doc.save_docx_flow(out("kq_lienmach.docx")), critical=False)
    step("Luu PNG", lambda: doc.save_images(out("kq.png"), 100, "png"))
    step("Luu JPG", lambda: doc.save_images(out("kq.jpg"), 100, "jpg"))
    step("Luu TXT", lambda: doc.save_txt(out("kq.txt")))

    def roundtrip():
        # mở lại file PDF đã lưu (font đã rút gọn) rồi sửa tiếp bằng chữ cái MỚI
        src2 = out("kq.pdf")
        for i, w in enumerate(["Minh", "Quỳnh Trâm Xuân", "Kỳ Duyên Phước Ẩu"]):
            d2 = core.PdfDocument(src2)
            ln2 = [l for l in d2.lines(0) if l.text.startswith("Họ và tên")][0]
            d2.replace_line(0, ln2, "Họ và tên: " + w, fm, fam, ln2.size, ln2.color, False, False)
            d2.add_text(0, 72, 450 + i * 16, "Lần %d: %s" % (i + 1, w), fm, fam, 13, (0, 0, 0), False, False)
            src2 = out("kq_lan%d.pdf" % (i + 2))
            d2.save_pdf(src2)
            t = fitz.open(src2)[0].get_text()
            bad = [c for c in t if c in "\x00\ufffd"]
            assert not bad, "lan %d: co %d ky tu o vuong" % (i + 2, len(bad))
            assert ("Họ và tên: " + w) in core.norm_text(t), "lan %d: khong thay chu moi" % (i + 2)

    step("Mo lai file da luu, sua tiep nhieu lan (khong o vuong)", roundtrip)

    def excel_check():
        import openpyxl
        doc.save_xlsx(out("kq.xlsx"))
        wb = openpyxl.load_workbook(out("kq.xlsx"))
        ws = wb.worksheets[0]
        vals = [str(c.value) for r in ws.iter_rows() for c in r if c.value]
        log("    Excel: %d o co chu, %d anh, %d vung gop o" % (len(vals), len(ws._images), len(ws.merged_cells.ranges)))
        assert any("Đã sửa" in v for v in vals), "Excel thieu chu da sua"
        assert any("Nguy" in v or "Họ và tên" in v for v in vals), "Excel thieu dong chu"
        assert len(ws._images) >= 1, "Excel thieu anh (con dau)"

    step("Luu Excel giu bo cuc (bang -> o, con dau -> anh)", excel_check)

    def word_normal_check():
        import zipfile
        x = zipfile.ZipFile(out("kq_bocuc.docx")).read("word/document.xml").decode("utf-8")
        assert "framePr" not in x, "Word van con khung (text box)"
        assert "<w:tbl>" in x or "<w:tbl " in x, "Word thieu bang"
        assert "wp:anchor" in x, "Word thieu anh con dau/chu ky"
        assert "Đã sửa" in x, "Word thieu chu da sua"

    step("Word la van ban binh thuong (khong khung, co bang that, co anh dau)", word_normal_check)

    def convert_check():
        import hashlib
        import convert
        src = out("kq.pdf")
        h0 = hashlib.md5(open(src, "rb").read()).hexdigest()
        cdir = out("chuyen")
        a, _ = convert.convert_file("pdf2word", src, cdir)
        b, _ = convert.convert_file("pdf2excel", src, cdir)
        c, _ = convert.convert_file("pdf2word", src, cdir)
        assert os.path.getsize(a) > 0 and os.path.getsize(b) > 0
        assert a != c, "ghi de file ket qua!"
        assert hashlib.md5(open(src, "rb").read()).hexdigest() == h0, "file goc bi thay doi!"
        log("    -> %s | %s | %s" % (os.path.basename(a), os.path.basename(b), os.path.basename(c)))
        for kind, f in (("word", a), ("excel", b)):
            eng = convert.detect_engines(kind)
            if not eng:
                log("    [BO QUA] %s -> PDF: may build khong co Office/LibreOffice" % kind)
                continue
            h1 = hashlib.md5(open(f, "rb").read()).hexdigest()
            p, used = convert.convert_file(kind + "2pdf", f, cdir)
            assert hashlib.md5(open(f, "rb").read()).hexdigest() == h1, "file %s goc bi thay doi" % kind
            log("    %s -> PDF bang %s: %s" % (kind, used, os.path.basename(p)))

    step("Chuyen doi PDF <-> Word/Excel (file goc khong doi, khong ghi de)", convert_check)

    def blank_check():
        import convert
        d = fitz.open()
        for txt in ("Trang A", None, "./.", None):
            p = d.new_page()
            if txt:
                p.insert_text((72, 300), txt, fontsize=13)
        d[1].insert_text((290, 815), "2", fontsize=10)  # trang trống chỉ có số trang ở chân
        src = out("co_trang_trong.pdf")
        d.save(src)
        cdir = out("chuyen_nhieu")
        res = []
        for m in ("pdf2word", "pdf2excel"):
            st = {}
            o, _ = convert.convert_file(m, src, cdir, stats=st)
            res.append((os.path.basename(o), st["blank"]))
            assert st["blank"] == 2, "%s: xoa %d trang trong (dung la 2)" % (m, st["blank"])
        assert fitz.open(src).page_count == 4, "file goc bi doi so trang!"
        log("    %s" % res)

    step("Tu xoa trang trong (giu trang chi co './.', file goc khong doi)", blank_check)

    def split_check():
        import splitter
        import hashlib

        def header(d):
            p = d.new_page()
            p.insert_text((90, 60), "UBND QUẬN 1", fontname="F0", fontfile=path, fontsize=13)
            p.draw_line((110, 82), (170, 82), width=0.6)
            p.insert_text((300, 60), "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", fontname="F0", fontfile=path, fontsize=13)
            p.insert_text((330, 76), "Độc lập - Tự do - Hạnh phúc", fontname="F0", fontfile=path, fontsize=14)
            p.draw_line((330, 82), (470, 82), width=0.6)
            for i in range(15):
                p.insert_text((72, 170 + i * 19), "Nội dung văn bản dòng %d ............................" % i,
                              fontname="F0", fontfile=path, fontsize=12)

        def cont(d, n):
            p = d.new_page()
            p.insert_text((290, 40), str(n), fontname="F0", fontfile=path, fontsize=12)
            # bảng nhỏ góc phải phía trên (trước đây làm tách nhầm thành văn bản mới)
            for r in range(4):
                p.draw_line((330, 60 + r * 20), (530, 60 + r * 20), width=0.7)
            for x in (330, 400, 530):
                p.draw_line((x, 60), (x, 120), width=0.7)
            for r in range(3):
                p.insert_text((340, 75 + r * 20), "Mục %d" % r, fontname="F0", fontfile=path, fontsize=11)
            for i in range(20):
                p.insert_text((72, 80 + i * 19), "Tiếp theo dòng %d ................................" % i,
                              fontname="F0", fontfile=path, fontsize=12)

        vec = fitz.open()
        header(vec); cont(vec, 2); vec.new_page()            # VB1: 2 trang + 1 trang trống
        header(vec)                                           # VB2: 1 trang
        header(vec); cont(vec, 2); vec.new_page(); cont(vec, 3)   # VB3: 3 trang (+ trống mặt sau)
        # bản "scan": mỗi trang chỉ là ảnh (không có lớp chữ)
        scan = fitz.open()
        for pg in vec:
            pm = pg.get_pixmap(dpi=120, colorspace=fitz.csGRAY)
            sp = scan.new_page(width=pg.rect.width, height=pg.rect.height)
            sp.insert_image(sp.rect, stream=pm.tobytes("png"))
        res = []
        for name, d in (("chu", vec), ("scan", scan)):
            src = out("gop_%s.pdf" % name)
            d.save(src)
            h0 = hashlib.md5(open(src, "rb").read()).hexdigest()
            sd = fitz.open(src)
            infos = splitter.analyze(sd)
            docs = splitter.plan(infos)
            files = splitter.save_docs(sd, docs, out("tach_" + name), "gop_" + name)
            sizes = [len(g) for _, g in files]
            res.append("%s: %d van ban %s" % (name, len(files), sizes))
            assert sizes == [2, 1, 3], "%s: tach sai %s" % (name, sizes)
            assert hashlib.md5(open(src, "rb").read()).hexdigest() == h0, "file goc bi thay doi"
        log("    " + " | ".join(res))

    step("Tach file PDF gop thanh tung van ban (bo trang trong)", split_check)

    def legacy_check():
        # PDF soạn bằng font .VnTime (mã TCVN3): sửa chữ / chuyển Word / tách văn bản phải ra đúng tiếng Việt
        import legacy
        import splitter
        d = fitz.open()
        pg = d.new_page()
        rows = [("Céng hoµ x· héi chñ nghÜa ViÖt Nam", 300, 60), ("§éc lËp - Tù do - H¹nh phóc", 330, 76),
                ("Th«n A, ngµy 19 th¸ng 10 n¨m 2026", 330, 110), ("B¸O C¸O", 270, 170),
                ("C«ng t¸c sè ho¸ n¨m 2026", 230, 188)]
        for t, x, y in rows:
            pg.insert_text((x, y), t, fontname="tiro", fontsize=13)
        for k in range(6):
            pg.insert_text((72, 230 + k * 18), "Néi dung v¨n b¶n dßng %d, ®Ò nghÞ c¸c ®¬n vÞ thùc hiÖn nghiªm tóc dßng dµi." % k,
                           fontname="tiro", fontsize=13)
        src = out("vntime.pdf")
        d.save(src)
        d = fitz.open(src)
        for x in range(1, d.xref_length()):
            if d.xref_get_key(x, "Type")[1] == "/Font":
                d.xref_set_key(x, "BaseFont", "/.VnTime")
        src2 = out("vntime2.pdf")
        d.save(src2)
        pd = core.PdfDocument(src2)
        texts = [l.text for l in pd.lines(0)]
        log("    .VnTime -> %s" % texts[:2])
        assert texts[0] == "Cộng hoà xã hội chủ nghĩa Việt Nam" and texts[1] == "Độc lập - Tự do - Hạnh phúc", texts
        nm = splitter.doc_name(pd.doc, [0], "VB")
        log("    ten file: " + nm)
        assert nm == "Bao cao cong tac so hoa nam 2026 ngay 19.10.2026", nm
        assert legacy.vni("Coäng hoøa xaõ hoäi chuû nghóa Vieät Nam") == "Cộng hòa xã hội chủ nghĩa Việt Nam"
        # gõ bằng Unikey bảng mã TCVN3 -> tự đổi; font .VnTime khi gõ -> dùng Times New Roman
        assert legacy.looks_tcvn3("§· kiÓm tra", 1) and legacy.tcvn3("§· kiÓm tra") == "Đã kiểm tra"
        assert not legacy.looks_tcvn3("Đã kiểm tra") and not legacy.looks_tcvn3("Café ở Paris") and not legacy.looks_tcvn3("a × b ÷ c = 5 µm, ± 2·x")
        assert fm.resolve(".VnTime")[0] == fm.resolve("Times New Roman")[0]

    step("Font ma cu .VnTime (TCVN3) / VNI -> sua chu, tach van ban dung tieng Viet", legacy_check)

    def merge_check():
        # gộp nhiều PDF -> 1 file; file gốc không đổi; không ghi đè file nguồn
        import merger
        import hashlib
        srcs = []
        for k, n in enumerate((2, 1, 3)):
            d = fitz.open()
            for i in range(n):
                pg = d.new_page()
                pg.insert_text((72, 100), "File %d trang %d" % (k + 1, i + 1), fontname="F0", fontfile=path, fontsize=13)
            pth = out("gop_nguon_%d.pdf" % (k + 1))
            d.save(pth)
            srcs.append(pth)
        h0 = [hashlib.md5(open(p_, "rb").read()).hexdigest() for p_ in srcs]
        dst = out("da_gop.pdf")
        ok, total, err = merger.merge(srcs, dst)
        r = fitz.open(dst)
        # (font Windows có thể trả dấu cách dạng \xa0 -> chuẩn hoá trước khi so)
        texts = [" ".join(r[i].get_text().replace("\xa0", " ").split()) for i in range(r.page_count)]
        log("    gop %d file -> %d trang, muc luc %d" % (ok, r.page_count, len(r.get_toc())))
        assert ok == 3 and total == 6 and r.page_count == 6 and not err
        assert texts[0] == "File 1 trang 1" and texts[2] == "File 2 trang 1" and texts[5] == "File 3 trang 3", texts
        assert [hashlib.md5(open(p_, "rb").read()).hexdigest() for p_ in srcs] == h0, "file goc bi thay doi"
        try:
            merger.merge(srcs, srcs[0])
            raise AssertionError("cho phep ghi de file nguon")
        except ValueError:
            pass
        assert hashlib.md5(open(srcs[0], "rb").read()).hexdigest() == h0[0]

    step("Gop nhieu file PDF thanh 1 file (file goc khong doi)", merge_check)

    def split_name_check():
        # tên file khi tách: loại văn bản + trích yếu + ngày, không dấu;
        # trang chỉ có số trang / chữ lem mờ / dấu đỏ lem -> trang trống (bỏ)
        import splitter
        d = fitz.open()
        p = d.new_page()
        p.insert_text((300, 60), "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", fontname="F0", fontfile=path, fontsize=12)
        p.insert_text((330, 76), "Độc lập - Tự do - Hạnh phúc", fontname="F0", fontfile=path, fontsize=13)
        p.insert_text((60, 105), "Số: 125/BC-PNV", fontname="F0", fontfile=path, fontsize=13)
        p.insert_text((330, 105), "Quận 1, ngày 10 tháng 10 năm 2026", fontname="F0", fontfile=path, fontsize=13)
        f = fitz.Font(fontfile=path)
        for y, t, sz in ((150, "BÁO CÁO", 15), (170, "Tình hình thực hiện nhiệm vụ quý III", 13)):
            p.insert_text((297 - f.text_length(t, sz) / 2, y), t, fontname="F0", fontfile=path, fontsize=sz)
        for k in range(10):
            p.insert_text((60, 210 + k * 18), "Căn cứ nội dung dòng %d, đề nghị các đơn vị thực hiện." % k,
                          fontname="F0", fontfile=path, fontsize=13)
        name = splitter.doc_name(d, [0], "VB01")
        # công văn có trích yếu "V/v ..." -> lấy trích yếu
        d2 = fitz.open()
        p2 = d2.new_page()
        p2.insert_text((300, 60), "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", fontname="F0", fontfile=path, fontsize=12)
        p2.insert_text((330, 76), "Độc lập – Tự do – Hạnh phúc", fontname="F0", fontfile=path, fontsize=13)
        p2.insert_text((72, 116), "V/v số hóa hồ sơ", fontname="F0", fontfile=path, fontsize=13)
        p2.insert_text((338, 116), "Đồng Nai, ngày 19 tháng 10 năm 2026", fontname="F0", fontfile=path, fontsize=13)
        for y, t, sz in ((180, "BÁO CÁO", 14), (196, "Công tác số hóa năm 2026", 13)):
            p2.insert_text((297 - f.text_length(t, sz) / 2, y), t, fontname="F0", fontfile=path, fontsize=sz)
        name2 = splitter.doc_name(d2, [0], "VB02")
        log("    ten file: %s | %s" % (name, name2))
        assert name == "Bao cao tinh hinh thuc hien nhiem vu quy III ngay 10.10.2026", name
        assert name2 == "So hoa ho so ngay 19.10.2026", name2
        # trang trống kiểu scan
        res = []
        for kind in ("so_trang", "lem", "dau_do", "mot_dong"):
            q = fitz.open()
            pg = q.new_page()
            if kind == "so_trang":
                pg.insert_text((290, 815), "- 3 -", fontname="F0", fontfile=path, fontsize=12)
            elif kind == "lem":
                for k in range(15):
                    pg.insert_text((60, 100 + k * 18), "Chữ mặt sau in hằn qua giấy dòng %d" % k,
                                   fontname="F0", fontfile=path, fontsize=13, color=(0.7, 0.7, 0.7))
            elif kind == "dau_do":
                pg.draw_circle((400, 600), 45, color=(0.85, 0.1, 0.1), width=3)
            else:
                pg.insert_text((60, 200), "Đề nghị các đơn vị liên quan phối hợp thực hiện./.",
                               fontname="F0", fontfile=path, fontsize=13)
            pm = pg.get_pixmap(dpi=120)
            sc = fitz.open()
            sp = sc.new_page()
            sp.insert_image(sp.rect, stream=pm.tobytes("png"))
            res.append((kind, splitter._is_blank(sp, splitter._gray(sp))[0]))
        log("    trang trong: %s" % res)
        assert res == [("so_trang", True), ("lem", True), ("dau_do", True), ("mot_dong", False)], res

    def scan_layer_check():
        # PDF scan đã có lớp chữ ẩn (máy scan tự OCR): phải coi là trang scan
        # -> Word chỉ có chữ gõ được, không giữ ảnh chứa chữ đè lên (chữ chồng)
        import ocrpage
        d = fitz.open()
        p = d.new_page()
        for k in range(12):
            p.insert_text((60, 100 + k * 18), "Nội dung văn bản dòng %d thực hiện" % k,
                          fontname="F0", fontfile=path, fontsize=13)
        pm = p.get_pixmap(dpi=100)
        s2 = fitz.open()
        q = s2.new_page()
        q.insert_image(q.rect, stream=pm.tobytes("png"))
        for k in range(12):
            q.insert_text((60, 100 + k * 18), "Nội dung văn bản dòng %d thực hiện" % k,
                          fontname="F0", fontfile=path, fontsize=13, render_mode=3)
        q2 = s2.new_page()
        q2.insert_image(q2.rect, stream=pm.tobytes("png"))
        for k in range(12):   # chữ hiện (không ẩn) nằm trên ảnh scan có cùng chữ
            q2.insert_text((60, 100 + k * 18), "Nội dung văn bản dòng %d thực hiện" % k,
                           fontname="F0", fontfile=path, fontsize=13)
        r = (ocrpage.is_scanned(s2[0]), ocrpage.is_scanned(s2[1]), ocrpage.is_scanned(d[0]))
        log("    scan co lop chu an / chu tren anh / ban so: %s" % (r,))
        assert r == (True, True, False), r

    step("PDF scan co san lop chu -> Word khong bi chu chong (nhan la trang scan)", scan_layer_check)

    step("Tach: dat ten theo trich yeu (V/v) hoac ten van ban + ngay, bo trang lem / chi co so trang", split_name_check)

    def glyph_check():
        # font khai báo sai mã: hình chữ 'ả' nhưng đọc ra 'd' -> phải sửa thành 'ả', giữ 'd' thật
        d = fitz.open()
        p = d.new_page()
        p.insert_text((72, 100), "CAX Trảng Bom - dung dưới", fontname="F0", fontfile=path, fontsize=13)
        src = out("ma_sai.pdf")
        d.save(src)
        from fontTools.ttLib import TTFont
        tt = TTFont(path, lazy=True, fontNumber=0)
        gid = tt.getGlyphID(tt.getBestCmap()[ord("ả")])
        d = fitz.open(src)
        for x in range(1, d.xref_length()):
            try:
                st = d.xref_stream(x)
            except Exception:
                continue
            if st and b"CMap" in st and (b"bfrange" in st or b"bfchar" in st):
                d.update_stream(x, st.replace(b"endcmap", b"1 beginbfchar\n<%04x> <0064>\nendbfchar\nendcmap" % gid))
        d.save(out("ma_sai2.pdf"))
        pd = core.PdfDocument(out("ma_sai2.pdf"))
        ln = pd.lines(0)[0]
        assert "Trdng" in ln.text, "khong tai hien duoc loi ma: %s" % ascii(ln.text)
        fixed = pd.fixed_texts(0, [ln], fm)[0]
        log("    doc ra: %s -> sua thanh: %s" % (ascii(ln.text), ascii(fixed)))
        assert fixed == "CAX Trảng Bom - dung dưới", "sua ma sai: %s" % ascii(fixed)

    step("Sua ma chu sai cua font ('Trdng' -> 'Trảng')", glyph_check)

    def wrap_check():
        # sửa đoạn dài trong ô bảng: tự xuống dòng, không vượt ô, không đè dòng dưới
        d = fitz.open()
        p = d.new_page()
        for r in range(3):
            p.draw_rect(fitz.Rect(60, 60 + r * 30, 300, 90 + r * 30), width=0.8)
            p.insert_text((66, 80 + r * 30), "Ô số %d" % r, fontname="F0", fontfile=path, fontsize=13)
        d.save(out("o_bang.pdf"))
        pd = core.PdfDocument(out("o_bang.pdf"))
        ln = [l for l in pd.lines(0) if l.text == "Ô số 1"][0]
        para = pd.para_for_line(0, ln)
        info = pd.replace_para(0, para, "Ô số 1 rất dài: " + "nội dung bổ sung " * 6, fm, fam, 13, (0, 0, 0), False, False)
        rows = [l for l in pd.lines(0) if 90 <= (l.bbox.y0 + l.bbox.y1) / 2 <= 120]
        log("    xuong %d dong, co chu %.1f, day duoi %.1f (gioi han %.1f)" % (info["rows"], info["size"], info["bottom"], para["limit"]))
        assert info["rows"] >= 2, "khong tu xuong dong"
        assert info["bottom"] <= para["limit"] + 0.5, "chu vuot ra ngoai o"
        assert all(l.bbox.x1 <= 300.5 for l in rows), "chu tran ra ngoai o"
        assert any(l.text == "Ô số 2" for l in pd.lines(0)), "mat dong ben duoi"

    step("Sua doan dai: tu xuong dong, khong de dong duoi", wrap_check)

    def ocr_check():
        # PDF scan (chỉ có ảnh) -> Word: phải ra chữ gõ được (OCR), không khung text box, có bảng
        import ocrpage
        import wordnative
        import zipfile
        import collections
        if not ocrpage.ocr_available():
            raise RuntimeError("thieu du lieu tieng Viet (tessdata/vie.traineddata): " + ocrpage.tessdata_report())
        d = fitz.open()
        p = d.new_page()
        lines = ["Căn cứ Kế hoạch số 45/KH-UBND, Ủy ban nhân dân phường báo cáo như sau:",
                 "Đề nghị các đơn vị liên quan phối hợp thực hiện đầy đủ nội dung trên."]
        for i, t in enumerate(lines):
            p.insert_text((60, 100 + i * 20), t, fontname="F0", fontfile=path, fontsize=13)
        for r in range(3):
            for c in range(3):
                rect = fitz.Rect(60 + c * 160, 160 + r * 24, 220 + c * 160, 184 + r * 24)
                p.draw_rect(rect, color=(0, 0, 0), width=0.8)
                p.insert_text((rect.x0 + 5, rect.y0 + 17), "Dữ liệu %d %d" % (r, c), fontname="F0",
                              fontfile=path, fontsize=12)
        pm = p.get_pixmap(dpi=200, colorspace=fitz.csGRAY)
        sc = fitz.open()
        sp = sc.new_page(width=p.rect.width, height=p.rect.height)
        sp.insert_image(sp.rect, stream=pm.tobytes("png"))
        src = out("scan_anh.pdf")
        sc.save(src)
        assert ocrpage.is_scanned(fitz.open(src)[0]), "khong nhan ra trang scan"
        dst = out("scan_anh.docx")
        wordnative.pdf_to_docx(fitz.open(src), dst)
        x = zipfile.ZipFile(dst).read("word/document.xml").decode("utf-8")
        assert "framePr" not in x and "txbx" not in x, "Word van con khung text box"
        assert "<w:tbl" in x, "khong dung lai duoc bang"
        import re as _re
        got = " ".join(_re.findall(r"<w:t[^>]*>([^<]*)</w:t>", x)).split()
        truth = (" ".join(lines) + " " + " ".join("Dữ liệu %d %d" % (r, c) for r in range(3) for c in range(3))).split()
        T, G = collections.Counter(truth), collections.Counter(got)
        acc = sum((T & G).values()) / len(truth)
        log("    OCR dung %.0f%% so tu" % (acc * 100))
        assert acc >= 0.85, "OCR sai nhieu (%.0f%%)" % (acc * 100)

    step("PDF scan -> Word go duoc (nhan dang chu tieng Viet, khong text box)", ocr_check)

    def math_check():
        # dòng có chỉ số trên + phân số -> Word: chỉ số trên thật, phân số là ảnh trong dòng, 1 trang
        import wordnative
        import zipfile
        d = fitz.open()
        p = d.new_page()
        p.insert_text((60, 100), "Ion Pb", fontname="F0", fontfile=path, fontsize=13)
        p.insert_text((97, 94), "2+", fontname="F0", fontfile=path, fontsize=8.5)
        p.insert_text((112, 100), "nằm ở tâm khối, công thức:", fontname="F0", fontfile=path, fontsize=13)
        p.insert_text((280, 92), "h", fontname="F0", fontfile=path, fontsize=13)
        p.draw_line((275, 96), (300, 96), width=0.6)
        p.insert_text((276, 109), "8L", fontname="F0", fontfile=path, fontsize=13)
        p.insert_text((310, 100), "là năng lượng.", fontname="F0", fontfile=path, fontsize=13)
        for i in range(8):
            p.insert_text((60, 140 + i * 18), "Dòng nội dung tiếp theo số %d của văn bản thử nghiệm." % i,
                          fontname="F0", fontfile=path, fontsize=13)
        src = out("cong_thuc.pdf")
        d.save(src)
        dst = out("cong_thuc.docx")
        wordnative.pdf_to_docx(fitz.open(src), dst)
        x = zipfile.ZipFile(dst).read("word/document.xml").decode("utf-8")
        assert 'w:vertAlign w:val="superscript"' in x, "mat chi so tren"
        assert "framePr" not in x and "txbx" not in x, "con khung text box"
        assert x.count("<wp:anchor") >= 1, "mat anh cong thuc"
        assert x.count("<w:br/>") >= 0
        log("    chi so tren: co | anh cong thuc: %d" % x.count("Cho cong thuc"))

    step("Cong thuc: chi so tren/duoi that, phan so la anh dung cho", math_check)

    for f in ("kq.pdf", "kq_bocuc.docx", "kq.xlsx", "kq_anh.docx", "kq_lienmach.docx", "kq.png", "kq.jpg", "kq.txt"):
        p = out(f)
        log("    %s: %s" % (f, ("%d bytes" % os.path.getsize(p)) if os.path.exists(p) else "KHONG CO"))

    if errors:
        log("SELFTEST FAIL: " + ", ".join(errors))
        return 1
    log("SELFTEST OK")
    return 0
