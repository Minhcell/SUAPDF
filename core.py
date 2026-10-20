# -*- coding: utf-8 -*-
"""
Lõi xử lý PDF cho phần mềm Sửa PDF Offline.
- Sửa chữ tại chỗ: xoá đúng chữ cũ (giữ nguyên khung bảng, đường kẻ, hình ảnh, chữ ký)
  rồi ghi chữ mới vào đúng vị trí cũ -> văn bản không bị nhảy.
- Xuất PDF / Word (giữ bố cục) / Word (dạng chảy) / Word (ảnh trang) / PNG / JPG / TXT.
"""
import os
import re
import sys
import json
import hashlib

try:
    import pymupdf as fitz
except ImportError:  # bản PyMuPDF cũ
    import fitz

# ----------------------------------------------------------------------------
#  FONT
# ----------------------------------------------------------------------------
DEFAULT_FAMILY = "Times New Roman"
FALLBACK_FAMILIES = ["Times New Roman", "Liberation Serif", "DejaVu Serif", "Arial"]


def app_data_dir():
    if sys.platform.startswith("win"):
        base = os.environ.get("LOCALAPPDATA") or os.environ.get("APPDATA") or os.path.expanduser("~")
    else:
        base = os.path.join(os.path.expanduser("~"), ".cache")
    d = os.path.join(base, "SuaPDF")
    os.makedirs(d, exist_ok=True)
    return d


def font_dirs():
    if sys.platform.startswith("win"):
        win = os.environ.get("WINDIR", r"C:\Windows")
        dirs = [os.path.join(win, "Fonts")]
        la = os.environ.get("LOCALAPPDATA")
        if la:
            dirs.append(os.path.join(la, "Microsoft", "Windows", "Fonts"))
    elif sys.platform == "darwin":
        dirs = ["/System/Library/Fonts", "/Library/Fonts", os.path.expanduser("~/Library/Fonts")]
    else:
        dirs = ["/usr/share/fonts", "/usr/local/share/fonts", os.path.expanduser("~/.fonts"),
                os.path.expanduser("~/.local/share/fonts")]
    # thư mục fonts đặt cạnh file exe (người dùng có thể tự chép thêm font .ttf vào)
    exe_dir = os.path.dirname(os.path.abspath(sys.executable if getattr(sys, "frozen", False) else __file__))
    dirs.append(os.path.join(exe_dir, "fonts"))
    return [d for d in dirs if os.path.isdir(d)]


def _name(tt, ids):
    nt = tt["name"]
    for nid in ids:
        for rec in nt.names:
            if rec.nameID == nid and rec.platformID == 3 and rec.langID == 0x409:
                try:
                    return rec.toUnicode()
                except Exception:
                    pass
        for rec in nt.names:
            if rec.nameID == nid:
                try:
                    return rec.toUnicode()
                except Exception:
                    pass
    return None


def scan_fonts(progress=None):
    """Trả về dict: family -> {'regular','bold','italic','bolditalic','vn'}"""
    from fontTools.ttLib import TTFont
    cache_path = os.path.join(app_data_dir(), "fonts_cache.json")
    try:
        with open(cache_path, "r", encoding="utf-8") as f:
            cache = json.load(f)
    except Exception:
        cache = {}
    files = []
    for d in font_dirs():
        for root, _, names in os.walk(d):
            for n in names:
                if n.lower().endswith((".ttf", ".otf")):
                    files.append(os.path.join(root, n))
    new_cache = {}
    for i, path in enumerate(files):
        try:
            st = os.stat(path)
            key = "%s|%d|%d" % (path, st.st_size, int(st.st_mtime))
        except OSError:
            continue
        info = cache.get(key)
        if info is None:
            try:
                tt = TTFont(path, lazy=True, fontNumber=0)
                fam = _name(tt, [1])
                sub = (_name(tt, [2]) or "Regular").lower()
                cmap = tt.getBestCmap() or {}
                vn = all(c in cmap for c in (0x1EC7, 0x01A1, 0x01B0, 0x0111, 0x1EA0))
                mac = tt["head"].macStyle if "head" in tt else 0
                bold = bool(mac & 1) or "bold" in sub
                italic = bool(mac & 2) or "italic" in sub or "oblique" in sub
                tt.close()
                info = {"family": fam, "bold": bold, "italic": italic, "vn": vn}
            except Exception:
                info = {"family": None}
        new_cache[key] = info
        if progress and i % 25 == 0:
            progress(i, len(files))
    try:
        with open(cache_path, "w", encoding="utf-8") as f:
            json.dump(new_cache, f)
    except Exception:
        pass
    fams = {}
    for key, info in new_cache.items():
        fam = info.get("family")
        if not fam:
            continue
        path = key.split("|")[0]
        style = ("bold" if info["bold"] else "") + ("italic" if info["italic"] else "")
        style = style or "regular"
        e = fams.setdefault(fam, {"vn": False})
        e["vn"] = e["vn"] or info.get("vn", False)
        if style not in e:
            e[style] = path
    # family phải có ít nhất 1 file dùng được
    out = {}
    for fam, e in fams.items():
        if any(k in e for k in ("regular", "bold", "italic", "bolditalic")):
            if "regular" not in e:
                e["regular"] = e.get("bold") or e.get("italic") or e.get("bolditalic")
            out[fam] = e
    return out


class FontManager:
    def __init__(self, families):
        self.families = families
        self._fonts = {}

    def default_family(self):
        for f in FALLBACK_FAMILIES:
            if f in self.families:
                return f
        return sorted(self.families)[0] if self.families else None

    def resolve(self, family, bold=False, italic=False):
        """-> (path, fake_bold)"""
        try:
            import legacy
            if legacy.is_legacy_family(family):
                # font mã cũ (.VnTime / VNI-Times) không gõ được chữ Unicode -> dùng font Unicode tương ứng
                family = legacy.unicode_family(family)
        except Exception:
            pass
        e = self.families.get(family) or self.families.get(self.default_family())
        if e is None:
            return None, False
        want = ("bold" if bold else "") + ("italic" if italic else "") or "regular"
        if want in e:
            return e[want], False
        if bold and italic and "bold" in e:
            return e["bold"], False
        if italic and "italic" in e:
            return e["italic"], bold
        return e["regular"], bold

    def font(self, path):
        f = self._fonts.get(path)
        if f is None:
            f = fitz.Font(fontfile=path)
            self._fonts[path] = f
        return f


def font_tag(path, session=""):
    """Tên font ghi vào PDF. Mỗi lần mở file dùng một mã phiên riêng, để KHÔNG BAO GIỜ
    dùng lại font đã bị rút gọn (subset) trong file đã lưu trước đó – font rút gọn chỉ
    còn các chữ cái cũ, gõ chữ mới sẽ hiện ô vuông."""
    return "S" + session + hashlib.md5(path.encode("utf-8")).hexdigest()[:6]


# ----------------------------------------------------------------------------
#  DỮ LIỆU DÒNG CHỮ
# ----------------------------------------------------------------------------
class Line:
    __slots__ = ("bbox", "text", "spans", "size", "color", "bold", "italic",
                 "origin", "font", "ascender", "horizontal", "invisible", "ocr")

    def __init__(self, ln, rot=None):
        spans = [s for s in ln["spans"] if s["text"].strip() or len(ln["spans"]) == 1]
        if not spans:
            spans = ln["spans"]
        self.spans = spans
        self.bbox = fitz.Rect(ln["bbox"])
        self.text = norm_text("".join(s["text"] for s in ln["spans"])).rstrip()
        s0 = max(spans, key=lambda s: len(s["text"].strip()))
        self.size = round(s0["size"] * 2) / 2
        c = s0.get("color", 0)
        self.color = ((c >> 16 & 255) / 255, (c >> 8 & 255) / 255, (c & 255) / 255)
        fl = s0.get("flags", 0)
        fname = s0.get("font", "")
        self.bold = bool(fl & 16) or "bold" in fname.lower()
        self.italic = bool(fl & 2) or "italic" in fname.lower()
        # toạ độ HIỂN THỊ (đã tính góc xoay của trang)
        rot = rot or fitz.Identity
        self.origin = fitz.Point(ln["spans"][0]["origin"]) * rot
        self.bbox = fitz.Rect(ln["bbox"]) * rot
        self.font = fname
        self.ascender = s0.get("ascender", 0.89) or 0.89
        d = fitz.Point(ln.get("dir", (1, 0))) * fitz.Matrix(rot.a, rot.b, rot.c, rot.d, 0, 0)
        self.horizontal = abs(d.x - 1) < 1e-3 and abs(d.y) < 1e-3
        # chữ ẩn (lớp OCR của bản scan): chữ thấy được thực ra là ảnh
        self.invisible = all(s.get("alpha", 255) == 0 for s in spans)
        self.ocr = False      # chữ do phần mềm nhận dạng lại (lớp chữ gốc hỏng / bản scan)


# Một số font (vd Times New Roman của Windows) dùng chung 1 hình chữ cho "-" và
# gạch mềm U+00AD, nên khi đọc lại PDF dấu "-" bị trả về ký tự lạ (ẩn, không hiện).
_NORM = {0x00AD: "-", 0x2010: "-", 0x2011: "-", 0x2012: "-", 0x2212: "-",
         0x00A0: " ", 0x2007: " ", 0x202F: " ", 0xFEFF: None, 0x200B: None,
         0x0000: None, 0xFFFD: None}  # ký tự hỏng (ô vuông) -> bỏ, để gõ lại


def _legacy_fix(dd):
    """Chữ mã cũ (TCVN3 .VnTime / VNI-Times) -> Unicode."""
    try:
        import legacy
        legacy.fix_dict(dd)
    except Exception:
        pass


def norm_text(s):
    return s.translate(_NORM)


def clean_font_name(name):
    """'ABCDEF+TimesNewRomanPS-BoldMT' -> 'Times New Roman'"""
    n = name.split("+")[-1]
    n = re.split(r"[-,]", n)[0]
    n = re.sub(r"(PSMT|PS|MT)$", "", n)
    known = {"TimesNewRoman": "Times New Roman", "Times": "Times New Roman",
             "LiberationSerif": "Times New Roman", "ArialMT": "Arial",
             "CourierNew": "Courier New", "Helvetica": "Arial", "Tahoma": "Tahoma",
             "Calibri": "Calibri", "Cambria": "Cambria", "Verdana": "Verdana"}
    if n in known:
        return known[n]
    if " " not in n:
        n = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", n)
    return n or DEFAULT_FAMILY


# ----------------------------------------------------------------------------
#  TÀI LIỆU
# ----------------------------------------------------------------------------
REDACT_KW = dict(images=getattr(fitz, "PDF_REDACT_IMAGE_NONE", 0),
                 graphics=getattr(fitz, "PDF_REDACT_LINE_ART_NONE", 0),
                 text=getattr(fitz, "PDF_REDACT_TEXT_REMOVE", 0))


def apply_text_only_redactions(page):
    """Chỉ xoá CHỮ trong vùng đánh dấu – giữ hình, chữ ký, đường kẻ bảng."""
    try:
        page.apply_redactions(**REDACT_KW)
    except TypeError:  # PyMuPDF cũ không có tham số text/graphics
        page.apply_redactions(images=REDACT_KW["images"])


IMAGE_EXT = (".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".gif", ".webp", ".jfif")
OFFICE_EXT = (".doc", ".docx", ".docm", ".rtf", ".odt", ".xls", ".xlsx", ".xlsm", ".xlsb", ".ods")


def open_as_pdf_bytes(path):
    """Đọc file thành PDF (trong bộ nhớ) để sửa: PDF giữ nguyên; ảnh (JPG/PNG/TIF... ảnh scan) -> mỗi ảnh 1 trang;
    Word/Excel -> chuyển sang PDF (cần Word/Excel hoặc LibreOffice trên máy). File gốc KHÔNG bị sửa."""
    ext = os.path.splitext(path)[1].lower()
    if ext in IMAGE_EXT:
        src = fitz.open(path)
        out = fitz.open()
        for i in range(src.page_count):        # TIF nhiều trang -> nhiều trang
            one = fitz.open("pdf", src.convert_to_pdf(i, i))
            out.insert_pdf(one)
        return out.tobytes(garbage=3, deflate=True)
    if ext in OFFICE_EXT:
        import tempfile
        import convert
        tmp = os.path.join(tempfile.mkdtemp(prefix="suapdf_"), "tam.pdf")
        convert.office_to_pdf(path, tmp)
        with open(tmp, "rb") as f:
            return f.read()
    with open(path, "rb") as f:
        return f.read()


class PdfDocument:
    MAX_UNDO = 30

    def __init__(self, path, password=None):
        if isinstance(path, (list, tuple)):
            # mở NHIỀU file cùng lúc: ghép thành 1 tài liệu theo đúng thứ tự đã chọn (file gốc không đổi)
            paths = list(path)
            if len(paths) == 1:
                path = paths[0]
            else:
                out = fitz.open()
                for p in paths:
                    out.insert_pdf(fitz.open(stream=open_as_pdf_bytes(p), filetype="pdf"))
                base = os.path.splitext(os.path.basename(paths[0]))[0]
                self.path = os.path.join(os.path.dirname(paths[0]), base + "_gop_%d_file.pdf" % len(paths))
                self.sources = paths
                self.doc = fitz.open(stream=out.tobytes(garbage=3, deflate=True), filetype="pdf")
                self._init_state()
                return
        self.path = path
        self.sources = [path]
        self.doc = fitz.open(stream=open_as_pdf_bytes(path), filetype="pdf")
        self._init_state()

    def _init_state(self):
        self.undo_stack, self.redo_stack = [], []
        self.modified = False
        self._lines = {}
        self.session = os.urandom(3).hex()

    # ---- trạng thái
    def needs_pass(self):
        return self.doc.needs_pass

    def authenticate(self, pw):
        return self.doc.authenticate(pw)

    def page_count(self):
        return self.doc.page_count

    def has_digital_signature(self):
        try:
            for page in self.doc:
                for w in page.widgets() or []:
                    if w.field_type == fitz.PDF_WIDGET_TYPE_SIGNATURE:
                        return True
        except Exception:
            pass
        return False

    def snapshot(self):
        self.undo_stack.append(self.doc.tobytes())
        if len(self.undo_stack) > self.MAX_UNDO:
            self.undo_stack.pop(0)
        self.redo_stack.clear()
        self.modified = True

    def _reload(self, data):
        self.doc = fitz.open(stream=data, filetype="pdf")
        self._lines.clear()
        getattr(self, "_geo", {}).clear()

    def undo(self):
        if not self.undo_stack:
            return False
        self.redo_stack.append(self.doc.tobytes())
        self._reload(self.undo_stack.pop())
        self.modified = True
        return True

    def redo(self):
        if not self.redo_stack:
            return False
        self.undo_stack.append(self.doc.tobytes())
        self._reload(self.redo_stack.pop())
        self.modified = True
        return True

    # ---- đọc
    def lines(self, pno):
        if pno not in self._lines:
            page = self.doc[pno]
            flags = fitz.TEXT_PRESERVE_WHITESPACE | fitz.TEXT_PRESERVE_LIGATURES | getattr(fitz, "TEXT_MEDIABOX_CLIP", 0)
            d = page.get_text("dict", flags=flags)
            _legacy_fix(d)
            out = []
            for b in d["blocks"]:
                if b.get("type", 0) != 0:
                    continue
                for ln in b["lines"]:
                    if not "".join(s["text"] for s in ln["spans"]).strip():
                        continue
                    out.append(Line(ln, page.rotation_matrix))
            self._ocr_broken(pno, page, out)
            self._lines[pno] = out
        return self._lines[pno]

    def _ocr_broken(self, pno, page, lines):
        """Trang có lớp chữ HỎNG (font thiếu bảng mã, đọc ra ký tự lạ): nhận dạng lại chữ đang hiện trên
        trang (OCR tiếng Việt) để khi bấm sửa hiện đúng chữ tiếng Việt."""
        try:
            import legacy
            import ocrpage
            raw = ["".join(sp.get("text", "") for sp in l.spans) for l in lines]   # chưa bỏ ký tự hỏng
            broken = legacy.page_text_broken(raw)
            # bản SCAN có sẵn lớp chữ do máy scan nhận dạng (thường sai: "ceNG HoA xA ngI...") -> tự nhận dạng lại
            # bản SCAN chỉ có ảnh (không có lớp chữ) -> nhận dạng chữ trên ảnh để sửa trực tiếp
            scan = False
            if not broken:
                kinds = self.__dict__.setdefault("_scan_kind", {})
                if pno not in kinds:
                    kinds[pno] = (ocrpage.text_layer_over_scan(page) if lines else ocrpage.is_scanned(page)) \
                        and ocrpage.ocr_available()
                scan = kinds[pno]
            if not broken and not scan:
                return
            # chữ do chính phần mềm vừa ghi đè lên ảnh scan (đã đúng) -> không đụng tới
            ov = ocrpage.overlay_text_rects(page) if scan else []

            def is_mine(l):
                return any(l.bbox.contains(fitz.Point((r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2)) for r in ov)
            if scan:
                for l in lines:
                    if not is_mine(l):
                        l.invisible = True        # sửa chữ -> phủ trắng chữ trên ảnh scan rồi mới ghi chữ mới
            cache = self.__dict__.setdefault("_ocr_cache", {})
            if pno not in cache:
                cache[pno] = ocrpage.ocr_page_words(page, keep_text=not scan) if ocrpage.ocr_available() else []
            words = cache[pno]
            if scan:
                self._scan_lines(page, lines, words, is_mine)
                return
            # gán từng chữ nhận dạng được cho dòng gần nhất cùng hàng (khung chữ lớp gốc có thể hẹp hơn ảnh)
            todo = [i for i, (l, rt) in enumerate(zip(lines, raw))
                    if not is_mine(l) and (scan or legacy.bad_text(rt) or not l.text.strip())]
            got = {i: [] for i in todo}
            for w in words:
                cy = (w[1] + w[3]) / 2
                h = max(4.0, w[3] - w[1])
                best, bd = None, None
                for i in todo:
                    b = lines[i].bbox
                    if not (b.y0 - h * 0.3 <= cy <= b.y1 + h * 0.3):
                        continue
                    d = max(0.0, b.x0 - w[2], w[0] - b.x1)
                    if bd is None or d < bd:
                        best, bd = i, d
                if best is not None and bd < h * 3:
                    got[best].append(w)
            for i in todo:
                l, rt = lines[i], raw[i]
                ws = sorted(got[i], key=lambda w: w[0])
                if ws and not legacy.bad_text(" ".join(w[4] for w in ws)):
                    l.text = " ".join(w[4] for w in ws)
                    l.ocr = True
                    if scan:
                        # bản scan: vị trí / cỡ chữ lấy theo chữ thật trên ảnh (lớp chữ máy scan hay lệch)
                        bb = fitz.Rect(ws[0][:4])
                        for w in ws:
                            bb |= fitz.Rect(w[:4])
                        hs = sorted(w[3] - w[1] for w in ws)
                        hh = hs[len(hs) // 2]
                        l.bbox = bb
                        l.size = max(6.0, round(hh / 1.1 * 2) / 2)
                        l.origin = fitz.Point(bb.x0, bb.y1 - hh * 0.22)
                        l.ascender = 0.89
                elif legacy.bad_text(rt):
                    l.text = ""
        except Exception:
            pass

    def _scan_lines(self, page, lines, words, is_mine):
        """Bản scan có lớp chữ máy scan: dựng lại các dòng theo chữ tự nhận dạng trên ảnh
        (đúng vị trí, cỡ chữ, nội dung); lớp chữ cũ của máy scan gắn vào dòng để xoá khi sửa."""
        keep = [l for l in lines if is_mine(l)]
        old = [l for l in lines if not is_mine(l)]
        rows = []
        for w in sorted(words, key=lambda w: ((w[1] + w[3]) / 2, w[0])):
            cy = (w[1] + w[3]) / 2
            h = max(4.0, w[3] - w[1])
            for r in rows:
                hh = max(h, r["h"])
                if abs(cy - r["cy"]) < 0.5 * hh and w[0] - r["x1"] < 2.5 * hh and w[2] > r["x0"] - 2.5 * hh:
                    r["w"].append(w)
                    r["x0"], r["x1"] = min(r["x0"], w[0]), max(r["x1"], w[2])
                    break
            else:
                rows.append({"cy": cy, "h": h, "x0": w[0], "x1": w[2], "w": [w]})
        new = []
        der = page.derotation_matrix
        try:
            import ocrpage
            white = ocrpage.overlay_fill_rects(page)     # chỗ chữ cũ đã bị phủ trắng khi sửa / tẩy
        except Exception:
            white = []
        prepared = []
        for r in rows:
            ws = sorted(r["w"], key=lambda w: w[0])
            ws = [w for w in ws if not any(q.contains(fitz.Point((w[0] + w[2]) / 2, (w[1] + w[3]) / 2)) for q in white)]
            if not ws:
                continue
            bb = fitz.Rect(ws[0][:4])
            for w in ws:
                bb |= fitz.Rect(w[:4])
            prepared.append((bb, ws))
        try:
            import ocrpage
            styles = ocrpage.row_styles(page, prepared)
        except Exception:
            styles = [(None, False)] * len(prepared)
        for (bb0, ws), (st_size, st_bold) in zip(prepared, styles):
            text = " ".join(w[4] for w in ws)
            bb = fitz.Rect(ws[0][:4])
            for w in ws:
                bb |= fitz.Rect(w[:4])
            hs = sorted(w[3] - w[1] for w in ws)
            hh = hs[len(hs) // 2]
            size = st_size or max(6.0, round(hh / 1.1 * 2) / 2)
            ex = fitz.Rect(bb.x0 - 2 * hh, bb.y0, bb.x1 + 2 * hh, bb.y1)
            under = [l for l in old if (min(l.bbox.y1, bb.y1) - max(l.bbox.y0, bb.y0)) > 0.3 * min(l.bbox.height, bb.height)
                     and l.bbox.x1 > ex.x0 and l.bbox.x0 < ex.x1]
            pb = bb * der
            span = {"text": text, "bbox": tuple(pb), "size": size, "font": "", "flags": 0, "color": 0,
                    "ascender": 0.89, "origin": tuple(fitz.Point(bb.x0, bb.y1 - hh * 0.22) * der)}
            spans = [span] + [sp for l in under for sp in l.spans]
            if under:
                s0 = under[0].spans[0]
                span["flags"], span["color"] = s0.get("flags", 0), s0.get("color", 0)
            ln = Line({"spans": spans, "bbox": tuple(pb), "dir": (1, 0)}, page.rotation_matrix)
            ln.text = text
            ln.bbox = bb
            ln.size = size
            ln.origin = fitz.Point(bb.x0, bb.y1 - hh * 0.22)
            ln.ascender = 0.89
            ln.invisible = True          # sửa -> phủ trắng chữ trên ảnh rồi mới ghi chữ mới
            ln.ocr = True
            ln.horizontal = True
            ln.bold = bool(st_bold)
            new.append(ln)
        lines[:] = sorted(keep + new, key=lambda l: (l.bbox.y0, l.bbox.x0))

    def line_at(self, pno, x, y, pad=1.5):
        best = None
        for ln in self.lines(pno):
            r = fitz.Rect(ln.bbox)
            r.x0 -= pad; r.y0 -= pad; r.x1 += pad; r.y1 += pad
            if r.contains(fitz.Point(x, y)):
                if best is None or ln.bbox.get_area() < best.bbox.get_area():
                    best = ln
        return best

    def render(self, pno, scale):
        page = self.doc[pno]
        return page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False, annots=True)

    # ---- sửa
    @staticmethod
    def _shrunk(r, page, k=0.28):
        r = fitz.Rect(r) * page.rotation_matrix      # sang toạ độ hiển thị
        h = r.height
        r = fitz.Rect(r.x0 + 0.2, r.y0 + h * k, r.x1 - 0.2, r.y1 - h * k)
        return r * page.derotation_matrix             # trả về toạ độ gốc

    def _erase_line(self, page, line):
        for s in line.spans:
            page.add_redact_annot(self._shrunk(s["bbox"], page), fill=False)
        apply_text_only_redactions(page)

    def write_text(self, page, fm, x, baseline, text, family, size, color, bold, italic, line_gap=1.2):
        path, fake_bold = fm.resolve(family, bold, italic)
        if path is None:
            raise RuntimeError("Không tìm thấy font: %s" % family)
        tag = font_tag(path, self.session)
        kw = dict(fontsize=size, fontname=tag, fontfile=path, color=color)
        if fake_bold:
            kw.update(render_mode=2, border_width=size * 0.035)
        for i, t in enumerate(text.split("\n")):
            if t == "":
                continue
            # x, baseline là toạ độ hiển thị -> đổi về toạ độ gốc của trang
            p = fitz.Point(x, baseline + i * size * line_gap) * page.derotation_matrix
            page.insert_text(p, t, rotate=page.rotation, **kw)

    def replace_line(self, pno, line, new_text, fm, family, size, color, bold, italic):
        self.snapshot()
        page = self.doc[pno]
        self._erase_line(page, line)
        if line.invisible:
            # bản scan có OCR: phủ trắng chữ ảnh cũ rồi mới ghi chữ mới
            r = fitz.Rect(line.bbox)
            r = fitz.Rect(r.x0 - 1, r.y0 - 1, r.x1 + 1, r.y1 + 1) * page.derotation_matrix
            page.draw_rect(r, color=None, fill=(1, 1, 1), overlay=True, width=0)
        if new_text.strip():
            self.write_text(page, fm, line.origin.x, line.origin.y, new_text,
                            family, size, color, bold, italic)
        self._lines.pop(pno, None)

    # ------------------------------------------------------------------ chữ đúng (sửa mã sai)
    def fixed_texts(self, pno, lines, fm):
        """Chữ của các dòng, đã sửa mã chữ sai của font (vd 'Trdng' -> 'Trảng')."""
        try:
            import glyphfix
            page = self.doc[pno]
            r = fitz.Rect(lines[0].bbox)
            for l in lines[1:]:
                r |= l.bbox
            r = fitz.Rect(r.x0 - 2, r.y0 - 2, r.x1 + 2, r.y1 + 2)
            raw = page.get_text("rawdict", clip=r * page.derotation_matrix,
                                flags=fitz.TEXT_PRESERVE_WHITESPACE | fitz.TEXT_PRESERVE_LIGATURES)
            _legacy_fix(raw)
            glyphfix.get_fixer(fm, id(self)).fix_page_spans(page, raw)
            rot = page.rotation_matrix
            got = []
            for b in raw.get("blocks", []):
                for ln in b.get("lines", []):
                    t = "".join(sp.get("text", "") for sp in ln["spans"])
                    got.append((fitz.Rect(ln["bbox"]) * rot, t))
            out = []
            for l in lines:
                if getattr(l, "ocr", False):
                    out.append(l.text)
                    continue
                parts = [(rb.x0, t) for rb, t in got
                         if l.bbox.contains(fitz.Point((rb.x0 + rb.x1) / 2, (rb.y0 + rb.y1) / 2))]
                parts.sort()
                t = norm_text("".join(p[1] for p in parts)).strip()
                try:
                    import legacy
                    if legacy.bad_text(t):      # lớp chữ hỏng -> dùng chữ đã nhận dạng lại
                        t = ""
                    if not t and legacy.bad_text(l.text):
                        # chỉ riêng dòng này hỏng (vd tiêu đề in hoa dùng font khác) -> đọc lại hình chữ của dòng
                        import ocrpage
                        if ocrpage.ocr_available():
                            o = ocrpage.ocr_line(page, l.bbox)
                            if o and not legacy.bad_text(o):
                                t = o
                except Exception:
                    pass
                out.append(t if t else l.text)
            return out
        except Exception:
            return [l.text for l in lines]

    # ------------------------------------------------------------------ đoạn văn (sửa như Word)
    def _page_geo(self, pno):
        """Ô bảng + đường kẻ ngang của trang (toạ độ hiển thị) – tính 1 lần."""
        if not hasattr(self, "_geo"):
            self._geo = {}
        g = self._geo.get(pno)
        if g is None:
            page = self.doc[pno]
            rot = page.rotation_matrix
            cells = []
            try:
                for t in page.find_tables().tables:
                    for c in t.cells:
                        if c:
                            cells.append(fitz.Rect(c) * rot)
            except Exception:
                pass
            hs, vs = [], []
            try:
                import wordnative
                hs, vs = wordnative._segments(page, rot)
            except Exception:
                pass
            g = {"cells": cells, "hs": hs, "vs": vs}
            self._geo[pno] = g
        return g

    def para_for_line(self, pno, line):
        """Đoạn văn chứa dòng được bấm: các dòng của đoạn, vùng được phép viết (trái/phải),
        giới hạn dưới (không đè lên chữ / đường kẻ bên dưới), căn lề, khoảng cách dòng."""
        page = self.doc[pno]
        H = page.rect.height
        L = [l for l in self.lines(pno) if l.horizontal]
        geo = self._page_geo(pno)

        def center(l):
            return fitz.Point((l.bbox.x0 + l.bbox.x1) / 2, (l.bbox.y0 + l.bbox.y1) / 2)

        cell = None
        for c in geo["cells"]:
            if c.contains(center(line)) and (cell is None or c.get_area() < cell.get_area()):
                cell = c
        if cell is not None:
            pool = [l for l in L if cell.contains(center(l))]
            regL, regR = cell.x0 + 1.5, cell.x1 - 1.5
        else:
            pool = [l for l in L if not any(c.contains(center(l)) for c in geo["cells"])]
            regL = min([l.bbox.x0 for l in pool] + [line.bbox.x0])
            # lề phải: bằng lề trái (trang thường cân 2 bên) hoặc dòng dài nhất
            regR = max([l.bbox.x1 for l in pool] + [line.bbox.x1, page.rect.width - regL])
            # ô kẻ khung (không nhận được là bảng): vạch đứng hai bên dòng là mép ô
            ly0, ly1 = line.bbox.y0, line.bbox.y1
            for (y0, y1, x, _w) in geo.get("vs", []):
                if y0 <= ly0 + 1 and y1 >= ly1 - 1:
                    if x >= line.bbox.x1 - 0.5:
                        regR = min(regR, x - 1.5)
                    elif x <= line.bbox.x0 + 0.5:
                        regL = max(regL, x + 1.5)
        # ghép dòng thành đoạn (dòng trước phải "đầy" mới nối, cùng lề trái, cùng khoảng cách dòng)
        pool = sorted(pool, key=lambda l: (round(l.origin.y, 1), l.bbox.x0))
        full = regR - max(6.0, 0.035 * (regR - regL))
        paras, cur = [], None
        for l in pool:
            if cur:
                last = cur[-1]
                gap = l.origin.y - last.origin.y
                # dòng sau thuộc cùng đoạn nếu từ đầu tiên của nó KHÔNG vừa chỗ trống cuối dòng trước
                # (Word phải ngắt dòng) – đúng cho cả đoạn căn đều lẫn căn trái
                fw = l.text.split(" ")[0] if l.text else ""
                fw_w = l.bbox.width * len(fw) / max(1, len(l.text)) + l.size * 0.25
                wrapped = last.bbox.x1 >= full or last.bbox.x1 + fw_w > regR - 1
                # dòng trước gần kín bề rộng và dòng sau không phải gạch đầu dòng / số thứ tự
                if not wrapped and last.bbox.x1 >= regL + 0.85 * (regR - regL):
                    wrapped = not re.match(r"^\s*([-+•*–]|\d+[.)]|[a-zđ][.)])\s", l.text)
                ok = (abs(l.size - last.size) < 0.6 and 0.95 * last.size <= gap <= 1.75 * last.size
                      and wrapped and l.bbox.x0 <= cur[0].bbox.x0 + 1.0
                      and l.bbox.x0 >= regL - 1.0
                      # cùng cột: không nối dòng cột bên kia (vd "PHÒNG GIÁO DỤC" dưới cột Quốc hiệu)
                      and l.bbox.x0 >= cur[0].bbox.x0 - max(40.0, 3.5 * last.size)
                      and min(l.bbox.x1, last.bbox.x1) - max(l.bbox.x0, last.bbox.x0)
                      >= 0.5 * min(l.bbox.width, last.bbox.width))
                if ok and len(cur) >= 2:
                    pitch = (cur[-1].origin.y - cur[0].origin.y) / (len(cur) - 1)
                    ok = abs(gap - pitch) < 1.2 and abs(l.bbox.x0 - cur[1].bbox.x0) < 1.5
                if ok:
                    cur.append(l)
                    continue
            cur = [l]
            paras.append(cur)
        plines = next((p for p in paras if line in p), [line])
        n = len(plines)
        left = plines[1].bbox.x0 if n > 1 else plines[0].bbox.x0
        first_indent = plines[0].bbox.x0 - left
        maxx1 = max(l.bbox.x1 for l in plines)
        top = min(l.bbox.y0 for l in plines)
        bot = max(l.bbox.y1 for l in plines)
        ids = set(id(l) for l in plines)
        others = [l for l in L if id(l) not in ids]
        # vật cản bên phải / trái trên cùng hàng (vd cột Quốc hiệu bên cạnh)
        if cell is None:
            right_obs = [o.bbox.x0 for o in others if o.bbox.y1 > top + 1 and o.bbox.y0 < bot - 1
                         and o.bbox.x0 >= maxx1 - 1]
            if right_obs:
                regR = min(regR, min(right_obs) - 6)
            left_obs = [o.bbox.x1 for o in others if o.bbox.y1 > top + 1 and o.bbox.y0 < bot - 1
                        and o.bbox.x1 <= left + 1]
            if left_obs:
                regL = max(regL, max(left_obs) + 6)
        regR = max(regR, maxx1)
        # căn lề
        regC = (regL + regR) / 2
        if n > 1:
            if all(abs(l.bbox.x1 - maxx1) < 2 for l in plines[:-1]) and maxx1 >= regR - 8:
                align = "justify"
            elif all(abs((l.bbox.x0 + l.bbox.x1) / 2 - (plines[0].bbox.x0 + plines[0].bbox.x1) / 2) < 2 for l in plines):
                align = "center"
            else:
                align = "left"
        else:
            c = (line.bbox.x0 + line.bbox.x1) / 2
            if abs(c - regC) < 3 and line.bbox.x0 - regL > 4:
                align = "center"
            elif abs(line.bbox.x1 - regR) < 2 and line.bbox.x0 - regL > 25:
                align = "right"
            else:
                align = "left"
        if align == "center":
            c = (plines[0].bbox.x0 + plines[0].bbox.x1) / 2
            half = min(c - regL, regR - c)
            x0, x1 = c - half, c + half
        elif align == "right":
            x0, x1 = regL, maxx1
        else:
            x0, x1 = left, regR
        # giới hạn dưới: dòng chữ / đường kẻ gần nhất bên dưới trong phạm vi ngang
        limit = (cell.y1 - 1.0) if cell is not None else H - 20
        for o in others:
            if o.bbox.y0 >= bot - 1 and o.bbox.x1 > x0 + 1 and o.bbox.x0 < x1 - 1:
                limit = min(limit, o.bbox.y0 - 0.5)
        for (a, b, y, _w) in geo["hs"]:
            if y >= bot - 0.5 and b > x0 + 2 and a < x1 - 2:
                limit = min(limit, y - 1.0)
        pitch = ((plines[-1].origin.y - plines[0].origin.y) / (n - 1)) if n > 1 else round(line.size * 1.2, 2)
        return {"lines": plines, "x0": x0, "x1": x1, "first_indent": first_indent, "align": align,
                "pitch": pitch, "baseline": plines[0].origin.y, "top": top, "bottom": bot,
                "limit": limit, "cell": cell, "size": plines[0].size}

    @staticmethod
    def _wrap(font, text, size, width, first_indent=0.0):
        """Ngắt dòng như Word: theo từ, Enter = xuống dòng cứng. -> [(chuỗi, là_dòng_cuối_đoạn)]"""
        out = []
        for hard in text.split("\n"):
            line, first = "", True

            def avail():
                return width - (first_indent if first else 0.0)

            for w in hard.split(" "):
                trial = (line + " " + w) if line else w
                if font.text_length(trial, size) <= avail():
                    line = trial
                    continue
                if line:
                    out.append((line, False))
                    first = False
                    line = ""
                # từ dài hơn cả dòng -> cắt theo ký tự
                while font.text_length(w, size) > avail() and len(w) > 1:
                    k = len(w)
                    while k > 1 and font.text_length(w[:k], size) > avail():
                        k -= 1
                    out.append((w[:k], False))
                    first = False
                    w = w[k:]
                line = w
            out.append((line, True))
        return out

    def write_para(self, page, fm, text, family, size, color, bold, italic, x0, x1, baseline,
                   pitch, align="left", first_indent=0.0, limit=None, min_scale=0.6, min_size=4.0):
        """Viết đoạn văn tự xuống dòng trong [x0, x1]; nếu vượt giới hạn dưới thì thu nhỏ chữ."""
        path, fake_bold = fm.resolve(family, bold, italic)
        if path is None:
            raise RuntimeError("Không tìm thấy font: %s" % family)
        font = fm.font(path)
        asc, desc = font.ascender, abs(font.descender)
        top = baseline - asc * size
        width = max(10.0, x1 - x0)
        s = size
        while True:
            rows = self._wrap(font, text, s, width, first_indent)
            p = pitch * s / size
            b0 = top + asc * s
            bottom = b0 + (len(rows) - 1) * p + desc * s
            # thu nhỏ dần cho tới khi vừa – KHÔNG BAO GIỜ đè lên dòng / ô bên dưới
            # (cho phép nhỏ hơn min_scale nếu nội dung quá dài so với chỗ trống, tối thiểu min_size pt)
            if limit is None or bottom <= limit or s <= min_size + 1e-6:
                break
            s = max(min_size, s - 0.25)
        tag = font_tag(path, self.session)
        kw = dict(fontsize=s, fontname=tag, fontfile=path, color=color)
        if fake_bold:
            kw.update(render_mode=2, border_width=s * 0.035)
        para_start = True
        for i, (t, last) in enumerate(rows):
            y = b0 + i * p
            ind = first_indent if para_start else 0.0
            if t.strip():
                tw = font.text_length(t, s)
                if align == "center":
                    x = x0 + (width - tw) / 2
                elif align == "right":
                    x = x1 - tw
                else:
                    x = x0 + ind
                if align == "justify" and not last and " " in t.strip():
                    words = t.split(" ")
                    gaps = len(words) - 1
                    wsum = sum(font.text_length(w, s) for w in words)
                    sp = (width - ind - wsum) / gaps
                    xx = x0 + ind
                    for w in words:
                        if w:
                            page.insert_text(fitz.Point(xx, y) * page.derotation_matrix, w,
                                             rotate=page.rotation, **kw)
                        xx += font.text_length(w, s) + sp
                else:
                    page.insert_text(fitz.Point(x, y) * page.derotation_matrix, t, rotate=page.rotation, **kw)
            para_start = last
        return {"rows": len(rows), "size": s, "bottom": bottom}

    def replace_para(self, pno, para, new_text, fm, family, size, color, bold, italic):
        """Thay cả đoạn văn (1 hay nhiều dòng) – tự xuống dòng, không đè dòng dưới."""
        self.snapshot()
        page = self.doc[pno]
        for ln in para["lines"]:
            for sp in ln.spans:
                page.add_redact_annot(self._shrunk(sp["bbox"], page), fill=False)
        apply_text_only_redactions(page)
        for ln in para["lines"]:
            if ln.invisible:
                r = fitz.Rect(ln.bbox)
                r = fitz.Rect(r.x0 - 1, r.y0 - 1, r.x1 + 1, r.y1 + 1) * page.derotation_matrix
                page.draw_rect(r, color=None, fill=(1, 1, 1), overlay=True, width=0)
        info = {"rows": 0, "size": size, "bottom": para["bottom"]}
        if new_text.strip():
            pitch = para["pitch"] * size / max(para["size"], 1)
            info = self.write_para(page, fm, new_text, family, size, color, bold, italic,
                                   para["x0"], para["x1"], para["baseline"], pitch, para["align"],
                                   para["first_indent"], para["limit"])
        self._lines.pop(pno, None)
        return info

    def add_text(self, pno, x, top, new_text, fm, family, size, color, bold, italic, width=None):
        if not new_text.strip():
            return
        self.snapshot()
        page = self.doc[pno]
        path, _ = fm.resolve(family, bold, italic)
        asc = fm.font(path).ascender if path else 0.89
        if width:
            self.write_para(page, fm, new_text, family, size, color, bold, italic, x, x + width,
                            top + asc * size, size * 1.2, "left")
        else:
            self.write_text(page, fm, x, top + asc * size, new_text, family, size, color, bold, italic)
        self._lines.pop(pno, None)

    def erase_text_in_rect(self, pno, rect):
        """Xoá chữ trong vùng – giữ khung bảng, hình, chữ ký."""
        self.snapshot()
        page = self.doc[pno]
        page.add_redact_annot(fitz.Rect(rect) * page.derotation_matrix, fill=False)
        apply_text_only_redactions(page)
        self._lines.pop(pno, None)

    def whiteout_rect(self, pno, rect, color=(1, 1, 1)):
        """Phủ màu trắng toàn bộ vùng (dùng cho PDF scan)."""
        self.snapshot()
        page = self.doc[pno]
        page.draw_rect(fitz.Rect(rect) * page.derotation_matrix, color=None, fill=color, overlay=True, width=0)
        self._lines.pop(pno, None)

    # ------------------------------------------------------------------------
    #  XUẤT FILE
    # ------------------------------------------------------------------------
    def final_bytes(self):
        """PDF cuối cùng, font đã thu gọn (subset) cho nhẹ."""
        data = self.doc.tobytes()
        try:
            tmp = fitz.open(stream=data, filetype="pdf")
            tmp.subset_fonts()
            data2 = tmp.tobytes(garbage=3, deflate=True)
            # kiểm tra file mới mở lại được và số trang khớp
            chk = fitz.open(stream=data2, filetype="pdf")
            if chk.page_count == self.doc.page_count:
                return data2
        except Exception:
            pass
        return self.doc.tobytes(garbage=3, deflate=True)

    def save_pdf(self, out):
        data = self.final_bytes()
        tmp = out + ".tmp"
        with open(tmp, "wb") as f:
            f.write(data)
        os.replace(tmp, out)

    def save_images(self, out, dpi=200, fmt="png", progress=None):
        base, _ = os.path.splitext(out)
        n = self.doc.page_count
        files = []
        for i, page in enumerate(self.doc):
            pm = page.get_pixmap(dpi=dpi, alpha=False)
            p = out if n == 1 else "%s_trang_%03d.%s" % (base, i + 1, fmt)
            if fmt == "jpg":
                pm.save(p, jpg_quality=92) if "jpg_quality" in pm.save.__code__.co_varnames else pm.save(p)
            else:
                pm.save(p)
            files.append(p)
            if progress:
                progress(i + 1, n)
        return files

    def save_txt(self, out):
        with open(out, "w", encoding="utf-8-sig") as f:
            for i, page in enumerate(self.doc):
                if i:
                    f.write("\n\n")
                dd = page.get_text("dict", sort=True)
                _legacy_fix(dd)
                f.write("\n".join("".join(sp["text"] for sp in ln["spans"]).rstrip()
                                   for b in dd["blocks"] if b.get("type", 0) == 0 for ln in b["lines"]))

    def save_docx_flow(self, out, progress=None):
        """Word dạng chảy (pdf2docx): sửa thoải mái, bảng thành bảng Word, có thể lệch nhẹ."""
        from pdf2docx import Converter
        tmp_pdf = os.path.join(app_data_dir(), "_tmp_convert.pdf")
        with open(tmp_pdf, "wb") as f:
            f.write(self.doc.tobytes())
        cv = Converter(tmp_pdf)
        try:
            cv.convert(out, multi_processing=False)
        finally:
            cv.close()
            try:
                os.remove(tmp_pdf)
            except OSError:
                pass

    def save_docx_image(self, out, dpi=200, progress=None):
        """Word: mỗi trang là 1 ảnh nét cao – giống bản PDF 100%, không sửa được chữ."""
        return _docx_build(self.doc, out, dpi, with_text=False, progress=progress)

    def save_docx_layout(self, out, dpi=200, progress=None):
        """Word giữ nguyên bố cục, sửa như Word bình thường:
        đoạn văn thường + bảng Word thật (không frame/text box); chữ ký, con dấu là ảnh đặt đúng chỗ."""
        import wordnative
        return wordnative.pdf_to_docx(self.doc, out, dpi, progress)

    def save_xlsx(self, out, dpi=200, progress=None):
        """Excel giữ bố cục: bảng thành ô Excel, chữ ký / con dấu là ảnh đặt đúng chỗ."""
        import convert
        return convert.pdf_to_excel(self.doc, out, dpi, progress)


# ----------------------------------------------------------------------------
#  WORD
# ----------------------------------------------------------------------------
def _xml_escape(s):
    return (s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            .replace('"', "&quot;"))


def _clean_xml_text(s):
    return "".join(ch for ch in s if ch in "\t\n" or ord(ch) >= 32)


def _docx_build(doc, out, dpi, with_text, progress=None):
    import io
    from docx import Document
    from docx.shared import Pt, Emu
    from docx.enum.section import WD_SECTION
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn

    d = Document()
    # font mặc định Times New Roman
    st = d.styles["Normal"]
    st.font.name = DEFAULT_FAMILY
    st.font.size = Pt(13)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = parse_xml('<w:rFonts %s/>' % nsdecls("w"))
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), DEFAULT_FAMILY)
    pf = st.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0

    body = d.element.body
    for p in list(body.findall(qn("w:p"))):
        body.remove(p)

    n = doc.page_count
    pic_id = 1
    for i in range(n):
        page = doc[i]
        W, H = page.rect.width, page.rect.height
        sec = d.sections[-1] if i == 0 else d.add_section(WD_SECTION.NEW_PAGE)
        sec.page_width, sec.page_height = Pt(W), Pt(H)
        sec.orientation = 1 if W > H else 0
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Pt(0)
        sec.header_distance = sec.footer_distance = Pt(0)
        sec.gutter = Pt(0)

        # ---- ảnh nền
        hlines = []
        if with_text:
            rot = page.rotation_matrix
            rot0 = fitz.Matrix(rot.a, rot.b, rot.c, rot.d, 0, 0)
            dd = page.get_text("dict", flags=fitz.TEXT_PRESERVE_WHITESPACE | fitz.TEXT_PRESERVE_LIGATURES)
            _legacy_fix(dd)
            for b in dd["blocks"]:
                if b.get("type", 0) != 0:
                    continue
                for ln in b["lines"]:
                    dv = fitz.Point(ln.get("dir", (1, 0))) * rot0
                    if all(s.get("alpha", 255) == 0 for s in ln["spans"]):
                        continue  # chữ ẩn OCR: đã có trong ảnh nền
                    if abs(dv.x - 1) < 1e-3 and "".join(s["text"] for s in ln["spans"]).strip():
                        hlines.append(ln)
            tmp = fitz.open()
            tmp.insert_pdf(doc, from_page=i, to_page=i)
            tp = tmp[0]
            # chỉ xoá chữ ngang (sẽ đặt lại bằng chữ Word); chữ xoay dọc giữ trong ảnh nền
            for ln in hlines:
                for s_ in ln["spans"]:
                    if s_["text"].strip():
                        tp.add_redact_annot(PdfDocument._shrunk(s_["bbox"], tp, 0.15), fill=False)
            apply_text_only_redactions(tp)
            pm = tp.get_pixmap(dpi=dpi, alpha=False)
        else:
            pm = page.get_pixmap(dpi=dpi, alpha=False)
        img = io.BytesIO(pm.tobytes("png"))

        par = d.add_paragraph()
        par.paragraph_format.space_after = Pt(0)
        run = par.add_run()
        run.font.size = Pt(1)
        run.add_picture(img, width=Emu(int(W * 12700)), height=Emu(int(H * 12700)))
        inline = run._r.find(".//" + qn("wp:inline"))
        graphic = inline.find(qn("a:graphic"))
        cx, cy = int(W * 12700), int(H * 12700)
        anchor = parse_xml(
            '<wp:anchor %s distT="0" distB="0" distL="0" distR="0" simplePos="0" '
            'relativeHeight="0" behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">'
            '<wp:simplePos x="0" y="0"/>'
            '<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
            '<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
            '<wp:extent cx="%d" cy="%d"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
            '<wp:wrapNone/><wp:docPr id="%d" name="NenTrang%d"/>'
            '<wp:cNvGraphicFramePr/></wp:anchor>' % (nsdecls("wp", "a", "pic", "r"), cx, cy, pic_id, i + 1))
        pic_id += 1
        anchor.append(graphic)
        inline.getparent().replace(inline, anchor)

        # ---- chữ đặt đúng toạ độ
        if with_text:
            for ln in hlines:
                    spans = [s for s in ln["spans"] if s["text"] != ""]
                    # bỏ khoảng trắng đầu dòng để vị trí x đúng chữ đầu tiên
                    while spans and not spans[0]["text"].strip():
                        spans.pop(0)
                    first = spans[0]
                    x = (fitz.Rect(first["bbox"]) * rot).x0
                    smax = max(spans, key=lambda s: s["size"])
                    asc = smax.get("ascender", 0.89) or 0.89
                    base_y = (fitz.Point(first["origin"]) * rot).y
                    top = base_y - asc * smax["size"]
                    width = (fitz.Rect(ln["bbox"]) * rot).x1 - x
                    fw = max(width * 1.12 + 6, 12)
                    fw = min(fw, max(W - x, width + 2))
                    runs = []
                    for k, s in enumerate(spans):
                        t = norm_text(s["text"])
                        if k == 0:
                            t = t.lstrip()
                        if k == len(spans) - 1:
                            t = t.rstrip()
                        if not t:
                            continue
                        c = s.get("color", 0)
                        fl = s.get("flags", 0)
                        fname = s.get("font", "")
                        fam = _xml_escape(clean_font_name(fname))
                        sz = int(round(s["size"] * 2))
                        rp = '<w:rFonts w:ascii="%s" w:hAnsi="%s" w:cs="%s" w:eastAsia="%s"/>' % (fam, fam, fam, fam)
                        if fl & 16 or "bold" in fname.lower():
                            rp += "<w:b/>"
                        if fl & 2 or "italic" in fname.lower():
                            rp += "<w:i/>"
                        rp += '<w:color w:val="%06X"/><w:sz w:val="%d"/><w:szCs w:val="%d"/>' % (c & 0xFFFFFF, sz, sz)
                        runs.append('<w:r><w:rPr>%s</w:rPr><w:t xml:space="preserve">%s</w:t></w:r>'
                                    % (rp, _xml_escape(_clean_xml_text(t))))
                    if not runs:
                        continue
                    p = parse_xml(
                        '<w:p %s><w:pPr>'
                        '<w:framePr w:w="%d" w:hRule="auto" w:wrap="around" w:vAnchor="page" w:hAnchor="page" w:x="%d" w:y="%d"/>'
                        '<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
                        '<w:ind w:left="0" w:right="0" w:firstLine="0"/><w:jc w:val="left"/>'
                        '</w:pPr>%s</w:p>' % (nsdecls("w"), int(fw * 20), int(round(x * 20)),
                                              int(round(top * 20)), "".join(runs)))
                    sect_p = body.find(qn("w:sectPr"))
                    sect_p.addprevious(p)
        if progress:
            progress(i + 1, n)
    d.save(out)
    return out
